import io
import zipfile

import docx.opc.exceptions
from docx import Document
from docxtpl import DocxTemplate
from jinja2.exceptions import TemplateError

from . import docx_export


class DocxTemplateValidationError(Exception):
    """Raised when an uploaded .docx isn't usable as a Report Profile
    template — either it isn't a real Word document, or docxtpl can't
    render it against a realistic (dummy) context. Surfaced as a form
    error at upload time, not left to fail at export time against a real
    engagement."""


# A .docx is a zip container -- the form-level size cap only bounds the
# *compressed* upload, so a small file with an extreme compression ratio
# (a zip bomb) could still expand to gigabytes once python-docx/docxtpl
# decompress it. Check the archive's own declared uncompressed size and
# member count before ever handing it to Document()/DocxTemplate().
MAX_DOCX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCX_MEMBER_COUNT = 2000


def _check_not_a_zip_bomb(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_DOCX_MEMBER_COUNT:
                raise DocxTemplateValidationError(
                    f"That file has too many internal parts ({len(infos)}) to be a normal "
                    "Word document."
                )
            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocxTemplateValidationError(
                    "That file expands to an unreasonable size once decompressed and was "
                    "rejected before opening it fully."
                )
    except zipfile.BadZipFile:
        # Not a zip at all -- Document() below gives the real "not a valid
        # Word document" error for this case, so just let it through here.
        return


def validate_docx_template(data: bytes) -> None:
    """Real-format check, mirroring branding_views._detect_image_content_type's
    "verify with the real library, not the browser content-type" approach."""
    _check_not_a_zip_bomb(data)
    try:
        Document(io.BytesIO(data))
    except (docx.opc.exceptions.PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise DocxTemplateValidationError("That file isn't a valid Word (.docx) document.") from exc


def dry_run_render(profile) -> None:
    """Renders the just-uploaded template against a cheap dummy context —
    catches a malformed {%p for %}/{%tr for %} loop or an unknown tag at
    upload time, not at export time against a real engagement's data."""
    try:
        tpl, context = docx_export.build_dummy_context(profile)
        tpl.render(context, jinja_env=docx_export.sandboxed_jinja_env())
    except TemplateError as exc:
        raise DocxTemplateValidationError(
            f"This template couldn't be rendered — check its {{{{ tags }}}} and any "
            f"{{%p for %}}/{{%tr for %}} loop syntax against the placeholder reference below: {exc}"
        ) from exc
    except Exception as exc:
        raise DocxTemplateValidationError(f"This template couldn't be rendered: {exc}") from exc
