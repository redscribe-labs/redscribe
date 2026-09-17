import json
import os
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse

from django.utils import timezone

from apps.accounts.models import Permission, Role
from apps.checklist.models import ChecklistItem, ChecklistRun
from apps.crypto.services import encrypt_bytes, generate_project_key, get_data_key, record_aad
from apps.engagements.models import Engagement, EngagementMembership
from apps.findings.models import ClassificationTag, ContentSectionDefinition, Finding, FindingSection, RetestRecord

from . import trends
from .assembly import _select_findings
from .models import ReportProfile, ReportSettings, ReportTextBlockDefinition
from .placeholders import build_placeholder_context, render_placeholders
from .tiptap_render import tiptap_to_html

User = get_user_model()
TEST_PASSWORD = "a-very-long-test-password-123!"


def doc_json(*parts) -> str:
    content = []
    for part in parts:
        text, marks = (part, []) if isinstance(part, str) else part
        node = {"type": "text", "text": text}
        if marks:
            node["marks"] = [{"type": m} if isinstance(m, str) else m for m in marks]
        content.append({"type": "paragraph", "content": [node]})
    return json.dumps({"type": "doc", "content": content})


def make_user(role, **kwargs):
    kwargs.setdefault("username", f"user-{role.lower()}-{os.urandom(4).hex()}")
    kwargs.setdefault("email", f"{kwargs['username']}@example.com")
    user = User.objects.create(
        role=Role.objects.get(slug=role.lower()), auth_type=User.AuthType.LOCAL, **kwargs
    )
    user.set_password(TEST_PASSWORD)
    user.save()
    return user


def login(client: Client, user) -> None:
    from django_otp import login as otp_login
    from django_otp.plugins.otp_totp.models import TOTPDevice

    client.force_login(user)
    device = TOTPDevice.objects.create(user=user, name="test", confirmed=True)
    session = client.session
    request = type("R", (), {"session": session, "user": user})()
    otp_login(request, device)
    session.save()


def set_finding_section(finding, project_key, slug: str, content: str) -> None:
    definition = ContentSectionDefinition.objects.get(slug=slug)
    FindingSection.objects.update_or_create(
        finding=finding, definition=definition,
        defaults={
            "content_ciphertext": encrypt_bytes(
                content.encode(), project_key, associated_data=record_aad("finding", finding.pk, slug),
            ),
        },
    )


_CONTENT_SECTION_FLAGS = {
    "vulnerability-description": {},
    "business-impact": {},
    "testing-summary": {},
    "technical-details": {"supports_image_upload": True, "is_import_target": True},
    "remediation-testing": {
        "supports_image_upload": True, "portal_visible": True,
        "include_in_remediation_report_only": True,
    },
    "recommendations": {},
    "references": {},
    "note": {"portal_visible": True},
}

_TEXT_BLOCK_SLUGS = [
    "summary_of_testing", "summary_of_findings", "summary_of_remediations", "conclusions",
    "information_gathering", "assessment_objectives", "scope_of_testing", "permitted_actions",
    "constraints", "assumptions", "requirements", "equipment_and_tooling", "vulnerability_rating",
]


def _heading(level, text):
    return {"type": "heading", "attrs": {"level": level}, "content": [{"type": "text", "text": text}]}


def _tag(slug):
    return {"type": "paragraph", "content": [{"type": "text", "text": f"{{{{ {slug} }}}}"}]}


def seed_report_fixtures():
    for i, (slug, extra) in enumerate(_CONTENT_SECTION_FLAGS.items(), start=1):
        ContentSectionDefinition.objects.get_or_create(
            slug=slug, defaults={"label": slug.replace("-", " ").capitalize(), "order": i * 10, **extra},
        )

    if not ReportProfile.objects.filter(is_default=True).exists():
        template = {
            "type": "doc",
            "content": [
                _tag("cover_page"),
                _tag("table_of_contents"),
                _heading(1, "Executive Summary"),
                _tag("summary_of_testing"),
                _tag("summary_of_findings"),
                _tag("breakdown_of_findings"),
                _tag("summary_of_remediations"),
                _tag("conclusions"),

                _heading(1, "Detailed Findings"),
                _tag("information_gathering"),
                _tag("finding_details"),
                _tag("observations"),

                _heading(1, "Testing Methodology"),
                _heading(2, "Scoping and requirements development"),
                _tag("assessment_objectives"),
                _tag("scope_of_testing"),
                _tag("permitted_actions"),
                _tag("constraints"),
                _tag("assumptions"),
                _tag("requirements"),
                _heading(2, "Assessment team"),
                _tag("assessment_team"),
                _heading(2, "Assessment outcomes"),
                _tag("equipment_and_tooling"),
                _tag("testing_phases"),
                _heading(2, "Vulnerability rating"),
                _tag("vulnerability_rating"),

                _heading(1, "Document control"),
                _tag("document_control"),
                _heading(1, "Disclaimers"),
                _tag("disclaimers"),
            ],
        }
        default_profile, _created = ReportProfile.objects.get_or_create(
            name="Default", defaults={"template": template, "block_defaults": {}, "is_default": True},
        )
        for slug in _TEXT_BLOCK_SLUGS:
            ReportTextBlockDefinition.objects.get_or_create(
                profile=default_profile, slug=slug,
                defaults={
                    "label": slug.replace("_", " ").capitalize(),
                    "include_in_remediation_report_only": slug == "summary_of_remediations",
                },
            )


def setUpModule():
    seed_report_fixtures()


def make_finding(engagement, created_by, **kwargs):
    tag = ClassificationTag.objects.first()
    key = get_data_key(engagement)
    vulnerability_description = kwargs.pop("vulnerability_description", doc_json("desc"))
    finding = Finding.objects.create(
        engagement=engagement,
        title=kwargs.pop("title", "Test finding"),
        severity=kwargs.pop("severity", Finding.Severity.HIGH),
        workflow_status=kwargs.pop("workflow_status", Finding.WorkflowStatus.DRAFT),
        created_by=created_by,
        **kwargs,
    )
    set_finding_section(finding, key, "vulnerability-description", vulnerability_description)
    set_finding_section(finding, key, "technical-details", doc_json("tech"))
    if tag:
        finding.classifications.add(tag)
    return finding


class TiptapToHtmlTests(TestCase):
    def test_blank_input(self):
        self.assertEqual(tiptap_to_html(""), "")
        self.assertEqual(tiptap_to_html(None), "")

    def test_invalid_json_returns_empty(self):
        self.assertEqual(tiptap_to_html("not json"), "")

    def test_paragraph_and_marks(self):
        html = tiptap_to_html(doc_json(("hello", ["bold"])))
        self.assertIn("<strong>hello</strong>", html)

    def test_heading_level_clamped(self):
        doc = json.dumps({"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 99}, "content": [{"type": "text", "text": "Hi"}]},
        ]})
        html = tiptap_to_html(doc)
        self.assertIn("<h3>Hi</h3>", html)

    def test_lead_paragraph_gets_report_lead_class(self):
        doc = json.dumps({"type": "doc", "content": [
            {"type": "paragraph", "attrs": {"lead": True}, "content": [{"type": "text", "text": "Intro"}]},
        ]})
        html = tiptap_to_html(doc)
        self.assertIn('<p class="report-lead">Intro</p>', html)

    def test_non_lead_paragraph_has_no_class(self):
        doc = json.dumps({"type": "doc", "content": [
            {"type": "paragraph", "attrs": {"lead": False}, "content": [{"type": "text", "text": "Body"}]},
        ]})
        html = tiptap_to_html(doc)
        self.assertIn("<p>Body</p>", html)

    def test_page_break_renders_as_marker_div(self):
        doc = json.dumps({"type": "doc", "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": "Before"}]},
            {"type": "pageBreak"},
            {"type": "paragraph", "content": [{"type": "text", "text": "After"}]},
        ]})
        html = tiptap_to_html(doc)
        self.assertIn('<p>Before</p><div class="report-page-break"></div><p>After</p>', html)

    def test_deeply_nested_content_does_not_crash(self):
        doc = {"type": "text", "text": "bottom"}
        for _ in range(200):
            doc = {"type": "blockquote", "content": [doc]}
        doc = {"type": "doc", "content": [doc]}

        html = tiptap_to_html(json.dumps(doc))
        self.assertEqual(html.count("<blockquote>"), 60)
        self.assertNotIn("bottom", html)

    def test_text_is_escaped(self):
        html = tiptap_to_html(doc_json("<script>alert(1)</script>"))
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)

    def test_link_allowed_protocol(self):
        doc = json.dumps({"type": "doc", "content": [{"type": "paragraph", "content": [
            {"type": "text", "text": "click", "marks": [{"type": "link", "attrs": {"href": "https://example.com"}}]},
        ]}]})
        html = tiptap_to_html(doc)
        self.assertIn('href="https://example.com"', html)

    def test_link_disallowed_protocol_dropped(self):
        doc = json.dumps({"type": "doc", "content": [{"type": "paragraph", "content": [
            {"type": "text", "text": "click", "marks": [{"type": "link", "attrs": {"href": "javascript:alert(1)"}}]},
        ]}]})
        html = tiptap_to_html(doc)
        self.assertNotIn("<a ", html)
        self.assertIn("click", html)

    def test_image_dropped_without_resolver(self):
        doc = json.dumps({"type": "doc", "content": [
            {"type": "image", "attrs": {"src": "https://evil.example.com/x.png"}},
        ]})
        html = tiptap_to_html(doc)
        self.assertNotIn("<img", html)

    def test_image_resolver_can_block_external_src(self):
        doc = json.dumps({"type": "doc", "content": [
            {"type": "image", "attrs": {"src": "http://169.254.169.254/latest/meta-data/"}},
        ]})
        html = tiptap_to_html(doc, image_resolver=lambda src: None)
        self.assertNotIn("169.254", html)
        self.assertNotIn("<img", html)

    def test_image_resolver_inlines_allowed_src(self):
        doc = json.dumps({"type": "doc", "content": [{"type": "image", "attrs": {"src": "/blobs/abc/"}}]})
        html = tiptap_to_html(doc, image_resolver=lambda src: "data:image/png;base64,AAAA")
        self.assertIn('src="data:image/png;base64,AAAA"', html)

    def test_text_transform_applied_to_text_only(self):
        html = tiptap_to_html(doc_json("hello {{ name }}"), text_transform=lambda t: t.replace("{{ name }}", "world"))
        self.assertIn("hello world", html)

    def test_unknown_node_type_does_not_crash(self):
        doc = json.dumps({"type": "doc", "content": [{"type": "totallyMadeUp", "content": [
            {"type": "text", "text": "still here"},
        ]}]})
        html = tiptap_to_html(doc)
        self.assertIn("still here", html)

    def test_highlight_mark_renders_as_mark_tag(self):
        html = tiptap_to_html(doc_json(("important", ["highlight"])))
        self.assertIn("<mark>important</mark>", html)

    def _code_block_doc(self):
        return json.dumps({"type": "doc", "content": [{"type": "codeBlock", "content": [
            {"type": "text", "text": "plain "},
            {"type": "text", "text": "bold", "marks": [{"type": "bold"}]},
            {"type": "text", "text": " "},
            {"type": "text", "text": "flagged", "marks": [{"type": "highlight"}]},
        ]}]})

    def test_marks_survive_inside_code_block(self):
        html = tiptap_to_html(self._code_block_doc())
        self.assertIn("<pre><code>plain <strong>bold</strong> <mark>flagged</mark></code></pre>", html)

    def _table_doc(self, rows):
        return json.dumps({"type": "doc", "content": [{"type": "table", "content": rows}]})

    def test_table_with_real_header_cells_unchanged(self):
        rows = [
            {"type": "tableRow", "content": [{"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Phase"}]}]}]},
            {"type": "tableRow", "content": [{"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Scoping"}]}]}]},
        ]
        html = tiptap_to_html(self._table_doc(rows))
        self.assertIn("<tr><th><p>Phase</p></th></tr>", html)
        self.assertIn("<tr><td><p>Scoping</p></td></tr>", html)

    def test_table_with_no_header_cells_treats_first_row_as_header(self):
        rows = [
            {"type": "tableRow", "content": [
                {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Phase"}]}]},
                {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Notes"}]}]},
            ]},
            {"type": "tableRow", "content": [
                {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Scoping"}]}]},
                {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "Done"}]}]},
            ]},
        ]
        html = tiptap_to_html(self._table_doc(rows))
        self.assertIn("<tr><th><p>Phase</p></th><th><p>Notes</p></th></tr>", html)
        self.assertIn("<tr><td><p>Scoping</p></td><td><p>Done</p></td></tr>", html)


class PlaceholderTests(TestCase):
    def test_known_token_substituted(self):
        self.assertEqual(render_placeholders("Client: {{ client_name }}", {"client_name": "Acme"}), "Client: Acme")

    def test_unknown_token_left_untouched(self):
        self.assertEqual(render_placeholders("{{ not_a_real_token }}", {}), "{{ not_a_real_token }}")

    def test_value_with_quotes_does_not_break_substitution(self):
        result = render_placeholders("Name: {{ client_name }}", {"client_name": 'Say "hi" {oddly}'})
        self.assertEqual(result, 'Name: Say "hi" {oddly}')

    def test_build_placeholder_context_handles_blank_dates(self):
        engagement = Engagement.objects.create(client_name="Acme")
        ctx = build_placeholder_context(engagement=engagement, user=None)
        self.assertEqual(ctx["start_date"], "")
        self.assertEqual(ctx["client_name"], "Acme")


class MarkdownExportHtmlBreakoutTests(TestCase):

    def test_html_to_markdown_escapes_decoded_script_tag(self):
        from .markdown_render import ImageCollector, html_to_markdown

        html = "<p>&lt;script&gt;alert(document.cookie)&lt;/script&gt;</p>"
        md = html_to_markdown(html, ImageCollector())
        self.assertNotIn("<script>", md)
        self.assertIn("\\<script\\>", md)

    def test_table_md_escapes_plain_string_cell(self):
        from .markdown_render import _render_table_md
        from .report_ir import TableBlock

        block = TableBlock(columns=["Title"], rows=[["<script>alert(1)</script>"]])
        md = _render_table_md(block)
        self.assertNotIn("<script>", md)
        self.assertIn("\\<script\\>", md)

    def test_list_block_items_are_escaped(self):
        from .markdown_render import _render_block
        from .report_ir import ListBlock

        block = ListBlock(items=["<script>alert(1)</script>"])
        md = _render_block(block, None, None)
        self.assertNotIn("<script>", md)
        self.assertIn("\\<script\\>", md)

    def test_link_href_scheme_still_enforced_after_html_roundtrip(self):
        from .markdown_render import ImageCollector, html_to_markdown

        html = '<p><a href="javascript:alert(1)">click</a></p>'
        md = html_to_markdown(html, ImageCollector())
        self.assertNotIn("javascript:", md)
        self.assertIn("click", md)

    def test_lead_paragraph_renders_as_bold_not_a_heading_in_markdown(self):
        from .markdown_render import ImageCollector, html_to_markdown

        md = html_to_markdown('<p class="report-lead">Intro copy</p>', ImageCollector())
        self.assertEqual(md.strip(), "**Intro copy**")

    def test_markdown_toc_link_matches_the_actual_heading_anchor(self):
        from .markdown_render import render_document_markdown
        from .report_ir import ReportDocument, ReportMeta, Section, TableOfContentsBlock, assign_numbers

        document = ReportDocument(
            meta=ReportMeta(
                project_id="PT-1", client_name="Acme", test_type_label="", classification_label="Confidential",
                report_date="2026-01-01",
            ),
            sections=[
                Section(key="toc", title="", numbered=False, blocks=[TableOfContentsBlock()]),
                Section(key="exec_summary", title="Executive Summary", toc_entry=True),
            ],
        )
        assign_numbers(document)
        md, _files = render_document_markdown(document)

        self.assertIn("(#1-executive-summary)", md)
        self.assertIn("# 1 Executive Summary", md)

    def test_markdown_toc_includes_second_level_entries(self):
        from .markdown_render import render_document_markdown
        from .report_ir import ReportDocument, ReportMeta, Section, TableOfContentsBlock, assign_numbers

        document = ReportDocument(
            meta=ReportMeta(
                project_id="PT-1", client_name="Acme", test_type_label="", classification_label="Confidential",
                report_date="2026-01-01",
            ),
            sections=[
                Section(key="toc", title="", numbered=False, blocks=[TableOfContentsBlock()]),
                Section(
                    key="detailed_findings", title="Detailed Findings", toc_entry=True,
                    children=[Section(key="finding:1", title="SQL Injection", toc_entry=True)],
                ),
            ],
        )
        assign_numbers(document)
        md, _files = render_document_markdown(document)

        self.assertIn("- [1.1 SQL Injection](#11-sql-injection)", md)

    def test_link_href_cannot_splice_a_second_markdown_link(self):
        from .markdown_render import ImageCollector, html_to_markdown

        html = '<p><a href="https://x.com)[y](javascript:alert(1)">click</a></p>'
        md = html_to_markdown(html, ImageCollector())
        self.assertNotIn("(javascript:alert(1))", md)
        self.assertIn("<https://x.com)[y](javascript:alert(1)>", md)


class MarkAndCodeBlockRenderTests(TestCase):

    def test_highlight_mark_in_paragraph_renders_as_double_equals(self):
        from .markdown_render import ImageCollector, html_to_markdown

        md = html_to_markdown("<p>this is <mark>important</mark></p>", ImageCollector())
        self.assertIn("this is ==important==", md)

    def test_marks_inside_code_block_render_as_plain_text_in_markdown(self):
        from .markdown_render import ImageCollector, html_to_markdown

        html = "<pre><code>plain <strong>bold</strong> <mark>flagged</mark></code></pre>"
        md = html_to_markdown(html, ImageCollector())
        self.assertIn("```\nplain bold flagged\n```", md)
        self.assertNotIn("**bold**", md)
        self.assertNotIn("==flagged==", md)


class AssignNumbersTests(TestCase):
    # Pure report_ir.assign_numbers logic, no DB needed — but TestCase
    # (not a plain unittest) to match this module's convention.

    def _doc(self, sections):
        from .report_ir import ReportDocument, ReportMeta

        meta = ReportMeta(project_id="", client_name="", test_type_label="", classification_label="", report_date="")
        return ReportDocument(meta=meta, sections=sections)

    def test_numbered_children_of_separate_unnumbered_wrappers_count_contiguously(self):
        # Two anonymous (numbered=False) wrapper sections — e.g. two
        # separate un-headed tag placements in a template with no real
        # headings — used to each reset assign_numbers' counter to 0, so
        # both wrappers' first numbered child came out "1" instead of the
        # whole level counting 1, 2, 3, 4 across both wrappers.
        from .report_ir import Section, assign_numbers

        wrapper1 = Section(key="w1", title="", numbered=False, children=[
            Section(key="a", title="A"), Section(key="b", title="B"),
        ])
        wrapper2 = Section(key="w2", title="", numbered=False, children=[
            Section(key="c", title="C"), Section(key="d", title="D"),
        ])
        document = self._doc([wrapper1, wrapper2])
        assign_numbers(document)
        self.assertEqual([c.number for c in wrapper1.children], ["1", "2"])
        self.assertEqual([c.number for c in wrapper2.children], ["3", "4"])

    def test_normal_nested_numbering_unaffected(self):
        from .report_ir import Section, assign_numbers

        a = Section(key="A", title="A", children=[Section(key="a1", title="a1"), Section(key="a2", title="a2")])
        b = Section(key="B", title="B", children=[Section(key="b1", title="b1")])
        document = self._doc([a, b])
        assign_numbers(document)
        self.assertEqual(a.number, "1")
        self.assertEqual([c.number for c in a.children], ["1.1", "1.2"])
        self.assertEqual(b.number, "2")
        self.assertEqual([c.number for c in b.children], ["2.1"])


class CssStyleTagEscapeTests(TestCase):

    def test_build_preview_css_neutralizes_style_tag_breakout_in_font(self):
        from . import ir_render
        from .report_ir import ReportMeta

        meta = ReportMeta(
            project_id="PT-1", client_name="Acme", test_type_label="", classification_label="Confidential",
            report_date="2026-01-01", body_font='</style><script>alert(1)</script>',
        )
        css = ir_render.build_preview_css(meta)
        self.assertNotIn("</style>", css)
        self.assertNotIn("<script>", css)

    def test_pdf_export_print_css_neutralizes_style_tag_breakout(self):
        from .pdf_export import _css_content_escape

        escaped = _css_content_escape('</style><script>alert(1)</script>')
        self.assertNotIn("</style>", escaped)
        self.assertNotIn("<script>", escaped)


class ReportProfileSettingsPermissionTests(TestCase):

    def setUp(self):
        self.profile = ReportProfile.objects.create(name="SMOKE TEST perm profile")
        self.url = reverse("report_profiles:edit", args=[self.profile.pk])

    def test_superadmin_can_view(self):
        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 200)

    def test_team_lead_cannot_view(self):
        client = Client()
        login(client, make_user(User.Role.TEAM_LEAD))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 403)

    def test_consultant_cannot_view(self):
        client = Client()
        login(client, make_user(User.Role.CONSULTANT))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 403)

    def test_superadmin_can_save(self):
        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        resp = client.post(self.url, {"cover_title": "My Report", "template": ""})
        self.assertEqual(resp.status_code, 302)
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.cover_title, "My Report")

    def test_body_font_rejects_html_breakout_attempt(self):
        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        payload = 'Arial", sans-serif; }</style><script>alert(document.cookie)</script><style>a{content:"'
        resp = client.post(self.url, {"cover_title": "My Report", "template": "", "body_font": payload})
        self.assertEqual(resp.status_code, 200)
        self.profile.refresh_from_db()
        self.assertNotEqual(self.profile.body_font, payload)

    def test_table_header_color_model_validator_rejects_bad_value_via_full_clean(self):
        self.profile.table_header_color = '</style><script>alert(1)</script>'
        with self.assertRaises(ValidationError):
            self.profile.full_clean()


def _make_test_png() -> bytes:
    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (10, 10), color=(255, 0, 0)).save(buf, format="PNG")
    return buf.getvalue()


class FirmBrandingViewTests(TestCase):
    def setUp(self):
        self.edit_url = reverse("branding:edit")
        self.logo_url = reverse("branding:logo")

    def test_superadmin_can_view_and_save_name(self):
        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        resp = client.get(self.edit_url)
        self.assertEqual(resp.status_code, 200)

        resp = client.post(self.edit_url, {"firm_name": "Acme Security"})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(ReportSettings.get_solo().firm_name, "Acme Security")

    def test_team_lead_cannot_view_or_save(self):
        client = Client()
        login(client, make_user(User.Role.TEAM_LEAD))
        self.assertEqual(client.get(self.edit_url).status_code, 403)
        resp = client.post(self.edit_url, {"firm_name": "Sneaky Co"})
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(ReportSettings.get_solo().firm_name, "")

    def test_team_lead_with_report_settings_manage_still_cannot_touch_branding(self):
        team_lead = make_user(User.Role.TEAM_LEAD)
        permission = Permission.objects.get(codename="report_settings.manage")
        team_lead.role.permissions.add(permission)
        client = Client()
        login(client, team_lead)
        resp = client.post(self.edit_url, {"firm_name": "Sneaky Co"})
        self.assertEqual(resp.status_code, 403)

    def test_upload_and_serve_logo(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        upload = SimpleUploadedFile("logo.png", _make_test_png(), content_type="image/png")
        resp = client.post(self.edit_url, {"firm_name": "Acme Security", "logo": upload})
        self.assertEqual(resp.status_code, 302)

        settings_obj = ReportSettings.get_solo()
        self.assertTrue(settings_obj.firm_logo)
        self.assertEqual(settings_obj.firm_logo_content_type, "image/png")

        logo_resp = client.get(self.logo_url)
        self.assertEqual(logo_resp.status_code, 200)
        self.assertEqual(logo_resp["Content-Type"], "image/png")
        self.assertEqual(bytes(logo_resp.content), _make_test_png())

    def test_logo_not_found_when_unset(self):
        client = Client()
        resp = client.get(self.logo_url)
        self.assertEqual(resp.status_code, 404)

    def test_logo_is_publicly_servable_without_login(self):
        settings_obj = ReportSettings.get_solo()
        settings_obj.firm_logo = _make_test_png()
        settings_obj.firm_logo_content_type = "image/png"
        settings_obj.save()

        client = Client()
        resp = client.get(self.logo_url)
        self.assertEqual(resp.status_code, 200)

    def test_remove_logo(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        upload = SimpleUploadedFile("logo.png", _make_test_png(), content_type="image/png")
        client.post(self.edit_url, {"firm_name": "Acme Security", "logo": upload})
        self.assertTrue(ReportSettings.get_solo().firm_logo)

        client.post(self.edit_url, {"firm_name": "Acme Security", "remove_logo": "on"})
        settings_obj = ReportSettings.get_solo()
        self.assertFalse(settings_obj.firm_logo)
        self.assertEqual(settings_obj.firm_logo_content_type, "")

    def test_oversized_logo_rejected(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from .forms import MAX_LOGO_BYTES

        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        oversized = SimpleUploadedFile("logo.png", _make_test_png() + b"\x00" * MAX_LOGO_BYTES, content_type="image/png")
        resp = client.post(self.edit_url, {"firm_name": "Acme Security", "logo": oversized})
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(ReportSettings.get_solo().firm_logo)

    def test_non_image_content_type_rejected(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        upload = SimpleUploadedFile("logo.txt", b"not an image", content_type="text/plain")
        resp = client.post(self.edit_url, {"firm_name": "Acme Security", "logo": upload})
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(ReportSettings.get_solo().firm_logo)


class DocumentControlAndTimelineTests(TestCase):

    def setUp(self):
        from apps.engagements.lifecycle import transition_engagement
        from apps.engagements.models import EngagementStatusHistory

        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        self.team_lead = make_user(User.Role.TEAM_LEAD)
        EngagementMembership.objects.create(user=self.author, engagement=self.engagement)
        EngagementStatusHistory.objects.create(engagement=self.engagement, status=Engagement.Status.IN_PROGRESS, changed_by=self.author)
        transition_engagement(self.engagement, Engagement.Status.IN_REVIEW, actor=self.team_lead)

        profile = ReportProfile.objects.get(is_default=True)
        profile.block_defaults = {
            **profile.block_defaults, "summary_of_remediations": doc_json("REMEDIATIONS MARKER"),
        }
        profile.save(update_fields=["block_defaults"])

    def _build_document(self, is_remediation_report=False):
        from types import SimpleNamespace

        from .assembly import build_report_document, seeded_content

        content = seeded_content(self.engagement, SimpleNamespace(content={}))
        config = SimpleNamespace(content=content, is_remediation_report=is_remediation_report)
        return build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )

    def test_document_control_commentary_renders_before_table(self):
        from types import SimpleNamespace

        from .assembly import build_report_document, sanitize_content
        from .report_ir import TableBlock

        content = sanitize_content({"document_control": {"notes": doc_json("Retention note")}})
        config = SimpleNamespace(content=content, is_remediation_report=False)
        document = build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )
        document_control_section = next(s for s in document.sections if s.title == "Document control")
        notes_index = next(
            i for i, b in enumerate(document_control_section.blocks)
            if hasattr(b, "html") and "Retention note" in b.html
        )
        table_index = next(
            i for i, b in enumerate(document_control_section.blocks) if isinstance(b, TableBlock)
        )
        self.assertLess(notes_index, table_index)

    def test_document_control_reflects_engagement_history(self):
        document = self._build_document()
        document_control_section = next(s for s in document.sections if s.title == "Document control")
        rows = document_control_section.blocks[0].rows
        stages = [r[0] for r in rows]
        self.assertIn("In Progress", stages)
        self.assertIn("In Review", stages)
        self.assertIn("Report exported", stages)
        self.assertEqual(rows[-1][0], "Report exported")

    def test_summary_of_remediations_excluded_when_not_remediation_report(self):
        document = self._build_document(is_remediation_report=False)
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        html = "".join(b.html for b in exec_summary.blocks if hasattr(b, "html"))
        self.assertNotIn("REMEDIATIONS MARKER", html)

    def test_summary_of_remediations_included_when_remediation_report(self):
        document = self._build_document(is_remediation_report=True)
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        html = "".join(b.html for b in exec_summary.blocks if hasattr(b, "html"))
        self.assertIn("REMEDIATIONS MARKER", html)


    def test_only_top_level_and_their_direct_children_are_toc_entries(self):
        document = self._build_document()
        top = next(s for s in document.sections if s.title == "Executive Summary")
        self.assertTrue(top.toc_entry)
        testing_methodology = next(s for s in document.sections if s.title == "Testing Methodology")
        self.assertTrue(testing_methodology.toc_entry)
        scoping = testing_methodology.children[0]
        self.assertTrue(scoping.toc_entry)
        detailed_findings = next(s for s in document.sections if s.title == "Detailed Findings")
        self.assertTrue(detailed_findings.toc_entry)


class BargraphAndTocRenderTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        make_finding(
            self.engagement, self.author, title="Only one", severity=Finding.Severity.CRITICAL,
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )

    def _build_document(self):
        from .assembly import build_report_document, get_draft_config

        config = get_draft_config(self.engagement)
        return build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )

    def test_bargraph_includes_every_severity_even_at_zero(self):
        from .assembly import _SEVERITY_ORDER
        from .report_ir import BarGraphBlock

        document = self._build_document()
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        bargraph = next(b for b in exec_summary.blocks if isinstance(b, BarGraphBlock))
        labels = [bar.label for bar in bargraph.series]
        for severity in _SEVERITY_ORDER:
            display = dict(Finding.Severity.choices)[severity]
            self.assertIn(f"{display} (Open)", labels)
        zero_bars = [bar for bar in bargraph.series if bar.label != "Critical (Open)"]
        self.assertTrue(all(bar.value == 0 for bar in zero_bars))

    def test_toc_does_not_include_individual_finding_titles(self):
        from .ir_render import render_toc

        document = self._build_document()
        html = render_toc(document)
        self.assertNotIn("Only one", html)

    def _build_document_with_breakdown(self, breakdown):
        from .assembly import build_report_document, get_draft_config, sanitize_content

        config = get_draft_config(self.engagement)
        config.content = {**sanitize_content(config.content), "breakdown": breakdown}
        config.save()
        return build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )

    def test_breakdown_disabled_produces_no_section(self):
        from .report_ir import BarGraphBlock

        document = self._build_document_with_breakdown({"enabled": False})
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        self.assertFalse(any(isinstance(b, BarGraphBlock) for b in exec_summary.blocks))
        self.assertFalse(any(c.key == "vulnerabilities_table" for c in exec_summary.children))

    def test_breakdown_table_only_omits_chart(self):
        from .report_ir import BarGraphBlock

        document = self._build_document_with_breakdown({"enabled": True, "show_chart": False, "show_table": True})
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        self.assertFalse(any(isinstance(b, BarGraphBlock) for b in exec_summary.blocks))
        self.assertTrue(any(c.key == "vulnerabilities_table" for c in exec_summary.children))

    def test_breakdown_chart_only_omits_table(self):
        from .report_ir import BarGraphBlock

        document = self._build_document_with_breakdown({"enabled": True, "show_chart": True, "show_table": False})
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        self.assertTrue(any(isinstance(b, BarGraphBlock) for b in exec_summary.blocks))
        self.assertFalse(any(c.key == "vulnerabilities_table" for c in exec_summary.children))

    def test_breakdown_severity_filter_excludes_bar_and_table_rows(self):
        from .report_ir import BarGraphBlock

        document = self._build_document_with_breakdown({
            "enabled": True, "show_chart": True, "show_table": True, "severities": ["HIGH"],
        })
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        bargraph = next(b for b in exec_summary.blocks if isinstance(b, BarGraphBlock))
        self.assertTrue(all("Critical" not in bar.label for bar in bargraph.series))
        vuln_table_section = next(c for c in exec_summary.children if c.key == "vulnerabilities_table")
        self.assertEqual(vuln_table_section.blocks[0].rows, [])

    def test_breakdown_intro_renders_above_chart(self):
        from .report_ir import BarGraphBlock, RichTextBlock

        document = self._build_document_with_breakdown({
            "enabled": True, "show_chart": True, "show_table": True,
            "intro": doc_json("Intro marker text"),
        })
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        chart_index = next(i for i, b in enumerate(exec_summary.blocks) if isinstance(b, BarGraphBlock))
        intro_index = next(
            i for i, b in enumerate(exec_summary.blocks)
            if isinstance(b, RichTextBlock) and "Intro marker text" in b.html
        )
        self.assertLess(intro_index, chart_index)

    def test_vulnerabilities_table_has_no_heading(self):
        document = self._build_document_with_breakdown({"enabled": True, "show_chart": True, "show_table": True})
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        vuln_table_section = next(c for c in exec_summary.children if c.key == "vulnerabilities_table")
        self.assertEqual(vuln_table_section.title, "")
        self.assertFalse(vuln_table_section.numbered)

    def test_breakdown_commentary_renders_between_chart_and_table(self):
        from .report_ir import BarGraphBlock, RichTextBlock

        document = self._build_document_with_breakdown({
            "enabled": True, "show_chart": True, "show_table": True,
            "notes": doc_json("Commentary marker text"),
        })
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        chart_index = next(i for i, b in enumerate(exec_summary.blocks) if isinstance(b, BarGraphBlock))
        notes_index = next(
            i for i, b in enumerate(exec_summary.blocks)
            if isinstance(b, RichTextBlock) and "Commentary marker text" in b.html
        )
        self.assertGreater(notes_index, chart_index)
        vuln_table_position = next(
            i for i, c in enumerate(exec_summary.children) if c.key == "vulnerabilities_table"
        )
        # Blocks render fully before children (see ir_render._render_section),
        # so appearing anywhere in .blocks after the chart already puts the
        # commentary before the (child-section) table in document order.
        self.assertIsNotNone(vuln_table_position)

    def test_breakdown_status_filter_excludes_open_bars(self):
        from .report_ir import BarGraphBlock

        document = self._build_document_with_breakdown({
            "enabled": True, "show_chart": True, "show_table": True, "statuses": ["CLOSED"],
        })
        exec_summary = next(s for s in document.sections if s.title == "Executive Summary")
        bargraph = next(b for b in exec_summary.blocks if isinstance(b, BarGraphBlock))
        self.assertFalse(any("(Open)" in bar.label for bar in bargraph.series))

    def test_toc_nests_level_two_entries(self):
        from .ir_render import render_toc

        document = self._build_document()
        html = render_toc(document)
        methodology_start = html.index("Testing Methodology")
        scoping_start = html.index("Scoping and requirements development")
        self.assertGreater(scoping_start, methodology_start)
        between = html[methodology_start:scoping_start]
        self.assertIn("<ul>", between)
        self.assertNotIn("</ul>", between)

    def test_pdf_toc_gets_page_numbers_both_levels(self):
        from .pdf_export import build_pdf

        document = self._build_document()
        pdf_bytes = build_pdf(document)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        from .pdf_export import _PRINT_CSS_TEMPLATE

        self.assertIn("target-counter(attr(href), page)", _PRINT_CSS_TEMPLATE)

    def test_bargraph_track_height_increased(self):
        from .ir_render import _PREVIEW_CSS_BASE

        self.assertIn("height: 17rem", _PREVIEW_CSS_BASE)

    def test_table_uses_horizontal_separators_not_full_grid(self):
        from .ir_render import _PREVIEW_CSS_BASE

        self.assertIn("border-bottom: 1px solid", _PREVIEW_CSS_BASE)
        self.assertNotIn("border: 1px solid #e2e8f0", _PREVIEW_CSS_BASE)

    def test_finding_grandchild_not_a_toc_entry(self):
        # Findings render as real H3 headings but are never TOC entries —
        # a TOC listing one line per finding doesn't scale to a report with
        # dozens of them. Same for their own subsections (Affects,
        # Vulnerability description, ...).
        document = self._build_document()
        detailed_findings = next(s for s in document.sections if s.title == "Detailed Findings")
        finding_section = next(
            c for c in detailed_findings.children if c.key.startswith("finding:") and c.key.count(":") == 1
        )
        self.assertFalse(finding_section.toc_entry)
        self.assertTrue(finding_section.children)
        self.assertFalse(finding_section.children[0].toc_entry)


class FindingMetadataClassificationTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)

    def test_one_row_per_taxonomy(self):
        from .assembly import _finding_metadata_table

        finding = make_finding(self.engagement, self.author)
        finding.classifications.clear()
        owasp = ClassificationTag.objects.get_or_create(taxonomy="OWASP Top 10 2025", value="A05:2025 – Injection")[0]
        mitre = ClassificationTag.objects.get_or_create(taxonomy="MITRE ATT&CK", value="TA0001 – Initial Access")[0]
        finding.classifications.add(owasp, mitre)

        table = _finding_metadata_table(finding, ReportProfile.objects.get(is_default=True))
        rows = {row[0]: row[1] for row in table.rows}
        self.assertEqual(rows["OWASP Top 10 2025"], "A05:2025 – Injection")
        self.assertEqual(rows["MITRE ATT&CK"], "TA0001 – Initial Access")
        self.assertNotIn("Classification", rows)

    def test_cve_id_row_present_and_falls_back_to_em_dash(self):
        from .assembly import _finding_metadata_table

        with_cve = make_finding(self.engagement, self.author, cve_id="CVE-2024-12345")
        without_cve = make_finding(self.engagement, self.author)

        profile = ReportProfile.objects.get(is_default=True)
        rows_with = {row[0]: row[1] for row in _finding_metadata_table(with_cve, profile).rows}
        rows_without = {row[0]: row[1] for row in _finding_metadata_table(without_cve, profile).rows}

        self.assertEqual(rows_with["CVE ID"], "CVE-2024-12345")
        self.assertEqual(rows_without["CVE ID"], "—")

    def test_severity_and_status_names_are_overridable_per_profile(self):
        from .assembly import _finding_metadata_table
        from .report_ir import StyledText

        finding = make_finding(
            self.engagement, self.author, severity=Finding.Severity.CRITICAL, status=Finding.Status.OPEN,
        )
        default_profile = ReportProfile.objects.get(is_default=True)
        default_rows = {row[0]: row[1] for row in _finding_metadata_table(finding, default_profile).rows}
        self.assertEqual(default_rows["Severity rating"], StyledText("Critical", background_rgba=default_rows["Severity rating"].background_rgba))
        self.assertEqual(default_rows["Status"], StyledText("Open", background_rgba=default_rows["Status"].background_rgba))

        localized_profile = ReportProfile.objects.create(
            name="French labels", labels={"severity_critical": "Critique", "status_open": "Ouverte"},
        )
        localized_rows = {row[0]: row[1] for row in _finding_metadata_table(finding, localized_profile).rows}
        self.assertEqual(localized_rows["Severity rating"].text, "Critique")
        self.assertEqual(localized_rows["Status"].text, "Ouverte")

    def test_affects_renders_right_after_business_impact(self):
        from .assembly import build_report_document, get_draft_config

        finding = make_finding(
            self.engagement, self.author, workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        key = get_data_key(self.engagement)
        set_finding_section(finding, key, "vulnerability-description", doc_json("desc"))
        set_finding_section(finding, key, "business-impact", doc_json("impact"))
        set_finding_section(finding, key, "testing-summary", doc_json("summary"))

        config = get_draft_config(self.engagement)
        document = build_report_document(
            config=config, engagement=self.engagement,
            project_key=key, user=self.author,
        )

        def find(sections, key):
            for s in sections:
                if s.key == key:
                    return s
                found = find(s.children, key)
                if found:
                    return found
            return None

        finding_section = find(document.sections, f"finding:{finding.id}")
        self.assertIsNotNone(finding_section)
        titles = [c.title for c in finding_section.children]
        self.assertEqual(
            titles[titles.index("Business impact"):titles.index("Business impact") + 2],
            ["Business impact", "Affects"],
        )
        self.assertLess(titles.index("Affects"), titles.index("Testing summary"))

    def test_no_classifications_shows_placeholder_row(self):
        from .assembly import _finding_metadata_table

        finding = make_finding(self.engagement, self.author)
        finding.classifications.clear()

        table = _finding_metadata_table(finding, ReportProfile.objects.get(is_default=True))
        rows = {row[0]: row[1] for row in table.rows}
        self.assertEqual(rows["Classification"], "—")


class ReportCoverBrandingTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)

    def _build_document(self):
        from .assembly import build_report_document, get_draft_config
        from .ir_render import render_cover_html

        config = get_draft_config(self.engagement)
        document = build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )
        return document, render_cover_html(document.meta)

    def test_no_branding_configured_shows_neither(self):
        _, cover_html = self._build_document()
        self.assertNotIn("report-cover-logo", cover_html)
        self.assertNotIn("report-cover-firm", cover_html)
        self.assertNotIn("RedScribe", cover_html)

    def test_cover_title_reaches_cover(self):
        profile = ReportProfile.objects.get(is_default=True)
        profile.cover_title = "Custom Security Assessment"
        profile.save()

        _, cover_html = self._build_document()
        self.assertIn("Custom Security Assessment", cover_html)

    def test_default_cover_title_shown_when_unset(self):
        profile = ReportProfile.objects.get(is_default=True)
        profile.cover_title = "Penetration Test Report"
        profile.save()

        _, cover_html = self._build_document()
        self.assertIn("Penetration Test Report", cover_html)

    def test_firm_name_reaches_cover(self):
        settings_obj = ReportSettings.get_solo()
        settings_obj.firm_name = "Acme Security"
        settings_obj.save()

        _, cover_html = self._build_document()
        self.assertIn("Acme Security", cover_html)
        self.assertIn("report-cover-firm", cover_html)

    def test_logo_takes_priority_over_name(self):
        settings_obj = ReportSettings.get_solo()
        settings_obj.firm_name = "Acme Security"
        settings_obj.firm_logo = _make_test_png()
        settings_obj.firm_logo_content_type = "image/png"
        settings_obj.save()

        document, cover_html = self._build_document()
        self.assertIn("report-cover-logo", cover_html)
        self.assertIn("data:image/png;base64,", cover_html)
        self.assertNotIn("report-cover-firm", cover_html)
        self.assertEqual(document.meta.firm_name, "Acme Security")


class FindingPageBreakTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        make_finding(self.engagement, self.author, title="First", workflow_status=Finding.WorkflowStatus.QA_APPROVED)
        make_finding(self.engagement, self.author, title="Second", workflow_status=Finding.WorkflowStatus.QA_APPROVED)

    def _finding_sections(self, document):
        found = []

        def walk(sections):
            for s in sections:
                if s.key.startswith("finding:") and s.key.count(":") == 1:
                    found.append(s)
                walk(s.children)

        walk(document.sections)
        return found

    def _build_document(self, toggle=None):
        from .assembly import build_report_document, get_draft_config

        config = get_draft_config(self.engagement)
        if toggle is not None:
            config.content.setdefault("toggles", {})["page_break_per_finding"] = toggle
            config.save()
        return build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )

    def test_page_break_off_by_default(self):
        document = self._build_document()
        sections = self._finding_sections(document)
        self.assertEqual(len(sections), 2)
        self.assertTrue(all(not s.page_break_before for s in sections))

    def test_page_break_on_when_toggled(self):
        document = self._build_document(toggle=True)
        sections = self._finding_sections(document)
        self.assertEqual(len(sections), 2)
        self.assertTrue(all(s.page_break_before for s in sections))

    def test_preview_html_carries_more_page_break_markers_when_toggled_on(self):
        from .ir_render import render_body_html

        on_html = render_body_html(self._build_document(toggle=True))
        off_html = render_body_html(self._build_document(toggle=False))
        self.assertEqual(on_html.count("report-page-break") - off_html.count("report-page-break"), 2)


class BrandingFallbackChainTests(TestCase):
    def test_defaults_to_redscribe(self):
        from .branding import get_branding

        b = get_branding()
        self.assertFalse(b["has_logo"])
        self.assertEqual(b["display_name"], "RedScribe")

    def test_firm_name_overrides_display_name(self):
        from .branding import get_branding

        settings_obj = ReportSettings.get_solo()
        settings_obj.firm_name = "Acme Security"
        settings_obj.save()
        self.assertEqual(get_branding()["display_name"], "Acme Security")

    def test_has_logo_reflects_stored_bytes(self):
        from .branding import get_branding

        settings_obj = ReportSettings.get_solo()
        settings_obj.firm_logo = _make_test_png()
        settings_obj.firm_logo_content_type = "image/png"
        settings_obj.save()
        self.assertTrue(get_branding()["has_logo"])


class ExportViewTests(TestCase):
    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        EngagementMembership.objects.create(user=self.author, engagement=self.engagement)
        make_finding(
            self.engagement, self.author, title="Approved one",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )

    def test_non_member_cannot_view_options(self):
        client = Client()
        login(client, make_user(User.Role.CONSULTANT))
        resp = client.get(reverse("reports:options", args=[self.engagement.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_member_can_view_options(self):
        client = Client()
        login(client, self.author)
        resp = client.get(reverse("reports:options", args=[self.engagement.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "1")

    def test_included_count_reflects_a_narrowed_finding_selection(self):
        from . import assembly

        config = assembly.get_draft_config(self.engagement)
        config.content = {"findings": {"included_ids": []}}
        config.save(update_fields=["content"])

        client = Client()
        login(client, self.author)
        resp = client.get(reverse("reports:options", args=[self.engagement.pk]))
        self.assertEqual(resp.context["included_count"], 0)

    def test_html_export(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "html"]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/html; charset=utf-8")
        self.assertIn(b"Approved one", resp.content)
        self.assertIn(b'attachment; filename="', resp["Content-Disposition"].encode())

    def test_export_creates_log_entry_with_comment(self):
        from .models import ReportExportLog

        client = Client()
        login(client, self.author)
        client.post(
            reverse("reports:download", args=[self.engagement.pk, "html"]),
            {"comment": "Delivered to client for review"},
        )
        log = ReportExportLog.objects.get(engagement=self.engagement)
        self.assertEqual(log.fmt, "html")
        self.assertEqual(log.exported_by, self.author)
        self.assertEqual(log.comment, "Delivered to client for review")

    def test_export_history_shown_on_options_page(self):
        from .models import ReportExportLog

        ReportExportLog.objects.create(
            engagement=self.engagement, fmt="pdf", comment="Final delivery", exported_by=self.author,
        )
        client = Client()
        login(client, self.author)
        resp = client.get(reverse("reports:options", args=[self.engagement.pk]))
        self.assertContains(resp, "Final delivery")
        self.assertContains(resp, "PDF")

    def test_failed_generation_does_not_log_export(self):
        from .models import ReportExportLog

        client = Client()
        login(client, self.author)
        client.post(reverse("reports:download", args=[self.engagement.pk, "not-a-real-format"]))
        self.assertFalse(ReportExportLog.objects.filter(engagement=self.engagement).exists())

    def test_html_export_matches_pdf_content_structure(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "html"]))
        body = resp.content.decode("utf-8")

        self.assertIn('class="report-cover-page pdf-front"', body)
        self.assertIn(self.engagement.client_name, body)
        self.assertIn("Executive Summary", body)
        self.assertIn("Testing Methodology", body)
        self.assertIn("Table of Contents", body)
        self.assertIn(".report-document", body)

    def test_pdf_export(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "pdf"]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_pdf_export_without_password_is_not_encrypted(self):
        from pypdf import PdfReader

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "pdf"]))
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(PdfReader(BytesIO(resp.content)).is_encrypted)

    def test_pdf_export_with_password_is_encrypted(self):
        from pypdf import PdfReader

        client = Client()
        login(client, self.author)
        resp = client.post(
            reverse("reports:download", args=[self.engagement.pk, "pdf"]),
            data={"pdf_password": "correct horse battery staple"},
        )
        self.assertEqual(resp.status_code, 200)

        reader = PdfReader(BytesIO(resp.content))
        self.assertTrue(reader.is_encrypted)
        self.assertEqual(PdfReader(BytesIO(resp.content)).decrypt("wrong password"), 0)
        self.assertNotEqual(reader.decrypt("correct horse battery staple"), 0)
        self.assertIn("Acme", reader.pages[0].extract_text())

    def test_encrypt_pdf_uses_aes256(self):
        from unittest.mock import patch

        from weasyprint import HTML

        from . import pdf_export

        minimal_pdf = HTML(string="<p>x</p>").write_pdf()
        with patch("pypdf.PdfWriter.encrypt") as mock_encrypt:
            pdf_export.encrypt_pdf(minimal_pdf, "pw")
        mock_encrypt.assert_called_once_with(user_password="pw", owner_password="pw", algorithm="AES-256")

    def test_pdf_password_ignored_for_other_formats(self):
        client = Client()
        login(client, self.author)
        resp = client.post(
            reverse("reports:download", args=[self.engagement.pk, "md"]),
            data={"pdf_password": "should-not-apply-here"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/zip")

    def test_unknown_format_rejected(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "exe"]))
        self.assertEqual(resp.status_code, 400)

    def test_get_export_download_not_allowed(self):
        client = Client()
        login(client, self.author)
        resp = client.get(reverse("reports:download", args=[self.engagement.pk, "pdf"]))
        self.assertEqual(resp.status_code, 405)

    def test_export_download_is_audit_logged(self):
        from apps.audit.models import AuditLogEntry

        client = Client()
        login(client, self.author)
        client.post(reverse("reports:download", args=[self.engagement.pk, "pdf"]))
        entry = AuditLogEntry.objects.filter(action="reports:download").latest("created_at")
        self.assertEqual(entry.actor, self.author)
        self.assertEqual(entry.engagement_id, str(self.engagement.pk))
        self.assertEqual(entry.status_code, 200)


class TrendsPermissionTests(TestCase):
    def setUp(self):
        self.url = reverse("trends:view")
        self.engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(self.engagement)

    def test_superadmin_can_view(self):
        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 200)

    def test_disabled_feature_flag_blocks_even_superadmin(self):
        from apps.feature_flags.models import FeatureFlags

        flags = FeatureFlags.get_solo()
        flags.recurrence_and_trends = False
        flags.save()

        client = Client()
        login(client, make_user(User.Role.SUPERADMIN))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 403)

    def test_team_lead_can_view_by_default(self):
        client = Client()
        login(client, make_user(User.Role.TEAM_LEAD))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 200)

    def test_consultant_cannot_view(self):
        client = Client()
        login(client, make_user(User.Role.CONSULTANT))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 403)

    def test_senior_cannot_view_by_default(self):
        client = Client()
        login(client, make_user(User.Role.SENIOR))
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 403)

    def test_role_management_can_grant_trends_permission_to_another_role(self):
        custom_role = Role.objects.create(name="Analyst")
        custom_role.permissions.set(Permission.objects.filter(codename="reports.trends"))
        custom_user = make_user(User.Role.CONSULTANT)
        custom_user.role = custom_role
        custom_user.save(update_fields=["role"])

        client = Client()
        login(client, custom_user)
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 200)

    def test_team_lead_limited_to_visible_engagements_superadmin_sees_all(self):
        limited_role = Role.objects.create(name="Scoped Analyst")
        limited_role.permissions.set(Permission.objects.filter(codename="reports.trends"))
        limited_user = make_user(User.Role.CONSULTANT)
        limited_user.role = limited_role
        limited_user.save(update_fields=["role"])

        other_engagement = Engagement.objects.create(client_name="Globex")
        generate_project_key(other_engagement)
        make_finding(other_engagement, make_user(User.Role.TEAM_LEAD), title="Only in Globex")

        client = Client()
        login(client, limited_user)
        resp = client.get(self.url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["severity_chart"]["columns"], [])

        superadmin_client = Client()
        login(superadmin_client, make_user(User.Role.SUPERADMIN))
        resp = superadmin_client.get(self.url)
        self.assertNotEqual(resp.context["severity_chart"]["columns"], [])


class TrendsComputationTests(TestCase):
    def setUp(self):
        self.author = make_user(User.Role.CONSULTANT)
        self.engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(self.engagement)

    def test_severity_over_time_groups_by_month_and_severity(self):
        f1 = make_finding(self.engagement, self.author, severity=Finding.Severity.CRITICAL)
        f2 = make_finding(self.engagement, self.author, severity=Finding.Severity.LOW)
        Finding.objects.filter(pk__in=[f1.pk, f2.pk]).update(created_at=timezone.now())

        rows = trends.severity_over_time(Engagement.objects.filter(pk=self.engagement.pk))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["counts"]["CRITICAL"], 1)
        self.assertEqual(rows[0]["counts"]["LOW"], 1)
        self.assertEqual(rows[0]["total"], 2)

    def test_severity_trend_chart_scales_columns_to_a_shared_max(self):
        from .trends_views import _severity_trend_chart

        chart = _severity_trend_chart([
            {"month": "m1", "label": "Jan 2026", "counts": {"CRITICAL": 3}, "total": 3},
            {"month": "m2", "label": "Feb 2026", "counts": {"CRITICAL": 1}, "total": 1},
        ])
        jan, feb = chart["columns"]
        self.assertEqual(jan["segments"][0]["height_pct"], 100.0)
        self.assertAlmostEqual(feb["segments"][0]["height_pct"], 33.33, places=1)

    def test_severity_trend_chart_legend_excludes_absent_severities(self):
        from .trends_views import _severity_trend_chart

        chart = _severity_trend_chart([
            {"month": "m1", "label": "Jan 2026", "counts": {"CRITICAL": 2}, "total": 2},
        ])
        self.assertEqual(len(chart["legend"]), 1)
        self.assertEqual(chart["legend"][0]["label"], "Critical")

    def test_severity_trend_chart_handles_empty_input(self):
        from .trends_views import _severity_trend_chart

        chart = _severity_trend_chart([])
        self.assertEqual(chart["columns"], [])
        self.assertEqual(chart["legend"], [])

    def test_mttr_bars_scale_relative_to_slowest_severity(self):
        from .trends_views import _with_mttr_bars

        mttr = _with_mttr_bars({
            "overall_days": 5, "sample_size": 2,
            "by_severity": [
                {"severity": "CRITICAL", "label": "Critical", "mean_days": 10.0, "sample_size": 1},
                {"severity": "LOW", "label": "Low", "mean_days": 2.0, "sample_size": 1},
            ],
        })
        critical, low = mttr["by_severity"]
        self.assertEqual(critical["bar_pct"], 100.0)
        self.assertEqual(low["bar_pct"], 20.0)
        self.assertIn("rgba(183,28,28,1)", critical["style"])

    def test_mttr_bars_handles_empty_by_severity(self):
        from .trends_views import _with_mttr_bars

        mttr = _with_mttr_bars({"overall_days": None, "sample_size": 0, "by_severity": []})
        self.assertEqual(mttr["by_severity"], [])

    def test_mean_time_to_remediate_uses_earliest_fixed_record(self):
        finding = make_finding(self.engagement, self.author, severity=Finding.Severity.HIGH)
        Finding.objects.filter(pk=finding.pk).update(created_at=timezone.now() - timezone.timedelta(days=10))
        finding.refresh_from_db()

        later = RetestRecord.objects.create(finding=finding, status=Finding.RetestStatus.FIXED, tested_by=self.author)
        RetestRecord.objects.filter(pk=later.pk).update(created_at=timezone.now())

        result = trends.mean_time_to_remediate(Engagement.objects.filter(pk=self.engagement.pk))
        self.assertEqual(result["sample_size"], 1)
        self.assertAlmostEqual(result["overall_days"], 10, delta=1)

    def test_mean_time_to_remediate_ignores_non_fixed_records(self):
        finding = make_finding(self.engagement, self.author)
        RetestRecord.objects.create(finding=finding, status=Finding.RetestStatus.NOT_FIXED, tested_by=self.author)

        result = trends.mean_time_to_remediate(Engagement.objects.filter(pk=self.engagement.pk))
        self.assertEqual(result["sample_size"], 0)
        self.assertIsNone(result["overall_days"])

    def test_repeat_findings_by_client_requires_multiple_engagements(self):
        tag = ClassificationTag.objects.filter(taxonomy="OWASP Top 10 2021").first()
        second_engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(second_engagement)

        f1 = make_finding(self.engagement, self.author)
        f1.classifications.set([tag])
        results = trends.repeat_findings_by_client(Engagement.objects.filter(client_name="Acme"))
        self.assertEqual(results, [])

        f2 = make_finding(second_engagement, self.author)
        f2.classifications.set([tag])
        results = trends.repeat_findings_by_client(Engagement.objects.filter(client_name="Acme"))
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["client_name"], "Acme")
        self.assertEqual(results[0]["engagement_count"], 2)
        self.assertEqual(results[0]["finding_count"], 2)


class SelectFindingsMalformedConfigTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        self.approved = make_finding(
            self.engagement, self.author, title="Approved",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        make_finding(self.engagement, self.author, title="Still draft")

    def test_findings_cfg_not_a_dict(self):
        result = _select_findings(self.engagement, "not a dict")
        self.assertEqual(result, [self.approved])

    def test_included_ids_not_a_list(self):
        result = _select_findings(self.engagement, {"included_ids": "not-a-list"})
        self.assertEqual(result, [self.approved])

    def test_included_ids_contains_non_string_entries(self):
        result = _select_findings(self.engagement, {"included_ids": [123, None, {"a": 1}, str(self.approved.id)]})
        self.assertEqual(result, [self.approved])

    def test_severities_not_a_list(self):
        result = _select_findings(self.engagement, {"severities": 42})
        self.assertEqual(result, [self.approved])

    def test_statuses_not_a_list(self):
        result = _select_findings(self.engagement, {"statuses": {"nope": True}})
        self.assertEqual(result, [self.approved])


class SanitizeContentTests(TestCase):

    def test_entirely_non_dict_content_returns_defaults(self):
        from .assembly import default_content, sanitize_content

        self.assertEqual(sanitize_content("garbage"), default_content())

    def test_report_configure_save_with_malformed_document_control_does_not_break_next_load(self):
        engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(engagement)
        author = make_user(User.Role.CONSULTANT)
        EngagementMembership.objects.create(user=author, engagement=engagement)

        client = Client()
        login(client, author)
        resp = client.post(
            reverse("reports:configure", args=[engagement.pk]),
            {"content_json": json.dumps({"document_control": "not a dict"})},
        )
        self.assertEqual(resp.status_code, 302)

        resp2 = client.get(reverse("reports:configure", args=[engagement.pk]))
        self.assertEqual(resp2.status_code, 200)

    def test_build_report_document_survives_malformed_toggles(self):
        from types import SimpleNamespace

        from .assembly import build_report_document, get_draft_config

        engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(engagement)
        author = make_user(User.Role.CONSULTANT)
        config = get_draft_config(engagement)
        config.content = {"toggles": "not a dict"}
        config.save(update_fields=["content"])

        document = build_report_document(
            config=SimpleNamespace(content=config.content, is_remediation_report=False),
            engagement=engagement, project_key=get_data_key(engagement), user=author,
        )
        self.assertIsNotNone(document)

    def test_build_report_document_survives_malformed_observation_and_phase_rows(self):
        from types import SimpleNamespace

        from .assembly import build_report_document, get_draft_config

        engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(engagement)
        author = make_user(User.Role.CONSULTANT)
        config = get_draft_config(engagement)
        config.content = {
            "observations": ["also not a dict"],
            "testing_phases": [42],
        }
        config.save(update_fields=["content"])

        document = build_report_document(
            config=SimpleNamespace(content=config.content, is_remediation_report=False),
            engagement=engagement, project_key=get_data_key(engagement), user=author,
        )
        self.assertIsNotNone(document)

    def test_finding_checkboxes_reflect_saved_narrow_selection(self):
        from apps.findings.models import Finding

        engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(engagement)
        author = make_user(User.Role.CONSULTANT)
        EngagementMembership.objects.create(user=author, engagement=engagement)
        kept = make_finding(engagement, author, title="Kept finding", workflow_status=Finding.WorkflowStatus.QA_APPROVED)
        excluded = make_finding(
            engagement, author, title="Excluded finding", workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )

        client = Client()
        login(client, author)
        client.post(
            reverse("reports:configure", args=[engagement.pk]),
            {"content_json": json.dumps({"findings": {"included_ids": [str(kept.id)]}})},
        )

        resp = client.get(reverse("reports:configure", args=[engagement.pk]))
        html = resp.content.decode()
        kept_pos = html.index(f'value="{kept.id}"')
        excluded_pos = html.index(f'value="{excluded.id}"')
        self.assertIn("checked", html[kept_pos:kept_pos + 60])
        self.assertNotIn("checked", html[excluded_pos:excluded_pos + 60])


class GetDraftConfigRaceTests(TransactionTestCase):

    def test_concurrent_first_access_creates_exactly_one_config(self):
        import threading

        from django.db import connection

        from .assembly import get_draft_config
        from .models import ReportConfig

        engagement = Engagement.objects.create(client_name="Acme")
        results = []
        barrier = threading.Barrier(2)

        def worker():
            barrier.wait()
            try:
                config = get_draft_config(engagement)
                results.append(config.pk)
            finally:
                connection.close()

        threads = [threading.Thread(target=worker) for _ in range(2)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0], results[1])
        self.assertEqual(ReportConfig.objects.filter(engagement=engagement, is_template=False).count(), 1)


class ReportProfileViewTests(TestCase):

    def setUp(self):
        from . import assembly

        self.assembly = assembly
        self.client = Client()
        self.user = make_user(User.Role.SUPERADMIN)
        login(self.client, self.user)
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-1")
        generate_project_key(self.engagement)
        EngagementMembership.objects.create(user=self.user, engagement=self.engagement)

    def test_edit_page_leaves_severity_colors_blank_by_default(self):
        # Regression: the edit form used to pre-fill every severity color
        # field with its built-in default as the submitted value, so saving
        # the page untouched silently locked in explicit colors for every
        # severity — "empty cell background color" never actually applied,
        # since no cell was ever really left blank. Defaults should only
        # show up as placeholder hints, never as values that get submitted.
        profile = ReportProfile.objects.get(is_default=True)
        resp = self.client.get(reverse("report_profiles:edit", args=[profile.pk]))
        self.assertEqual(resp.status_code, 200)
        form = resp.context["form"]
        self.assertEqual(form["severity_color__CRITICAL__open"].value(), "")
        self.assertIn("rgba", form["severity_color__CRITICAL__open"].field.widget.attrs["placeholder"])

    def test_saving_profile_untouched_keeps_severity_colors_blank(self):
        from . import assembly, ir_render
        from apps.crypto.services import get_data_key
        from apps.findings.models import Finding

        profile = ReportProfile.objects.create(
            name="SMOKE TEST empty cell profile", empty_cell_background_color="rgba(11,22,33,1)",
            template={"type": "doc", "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "{{ breakdown_of_findings }}"}]},
            ]},
        )
        resp = self.client.post(reverse("report_profiles:edit", args=[profile.pk]), {
            "template": json.dumps(profile.template),
            "cover_title": profile.cover_title,
            "classification_label": profile.classification_label,
            "finding_id_prefix": profile.finding_id_prefix,
            "body_font": profile.body_font,
            "monospace_font": profile.monospace_font,
            "bullet_character": profile.bullet_character,
            "table_header_color": profile.table_header_color,
            "empty_cell_background_color": profile.empty_cell_background_color,
        })
        self.assertEqual(resp.status_code, 302)
        profile.refresh_from_db()
        self.assertEqual(profile.severity_colors["HIGH"], {"open": "", "closed": ""})

        make_finding(
            self.engagement, self.user, severity=Finding.Severity.HIGH, status=Finding.Status.OPEN,
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = profile
        config.save()
        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        html = ir_render.render_document_html(document)
        self.assertIn("rgba(11,22,33,1)", html)

    def test_saving_profile_untouched_keeps_labels_at_shipped_defaults(self):
        # Same "blank means inherit" pattern as severity colors — an
        # untouched profile must render byte-identical to before the
        # Labels tab existed.
        from . import assembly, ir_render
        from apps.crypto.services import get_data_key
        from apps.findings.models import Finding

        profile = ReportProfile.objects.create(
            name="SMOKE TEST labels default profile",
            template={"type": "doc", "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "{{ finding_details }}"}]},
            ]},
        )
        resp = self.client.post(reverse("report_profiles:edit", args=[profile.pk]), {
            "template": json.dumps(profile.template),
            "cover_title": profile.cover_title,
            "classification_label": profile.classification_label,
            "finding_id_prefix": profile.finding_id_prefix,
            "body_font": profile.body_font,
            "monospace_font": profile.monospace_font,
            "bullet_character": profile.bullet_character,
            "table_header_color": profile.table_header_color,
            "empty_cell_background_color": profile.empty_cell_background_color,
        })
        self.assertEqual(resp.status_code, 302)
        profile.refresh_from_db()
        self.assertEqual(profile.labels, {})

        finding = make_finding(
            self.engagement, self.user, workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = profile
        config.save()
        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        html = ir_render.render_document_html(document)
        self.assertIn("Affects", html)
        self.assertIn("Severity rating", html)

    def test_custom_label_overrides_default_text(self):
        from . import assembly, ir_render
        from apps.crypto.services import get_data_key
        from apps.findings.models import Finding

        profile = ReportProfile.objects.create(
            name="SMOKE TEST custom label profile",
            template={"type": "doc", "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "{{ finding_details }}"}]},
            ]},
        )
        resp = self.client.post(reverse("report_profiles:edit", args=[profile.pk]), {
            "template": json.dumps(profile.template),
            "cover_title": profile.cover_title,
            "classification_label": profile.classification_label,
            "finding_id_prefix": profile.finding_id_prefix,
            "body_font": profile.body_font,
            "monospace_font": profile.monospace_font,
            "bullet_character": profile.bullet_character,
            "table_header_color": profile.table_header_color,
            "empty_cell_background_color": profile.empty_cell_background_color,
            "label__affects": "Impact area",
            "label__finding_severity_rating": "Risk rating",
        })
        self.assertEqual(resp.status_code, 302)
        profile.refresh_from_db()
        self.assertEqual(profile.labels["affects"], "Impact area")
        self.assertEqual(profile.labels["finding_severity_rating"], "Risk rating")
        # Untouched fields never got an explicit entry — still inherit.
        self.assertNotIn("finding_status", profile.labels)

        finding = make_finding(
            self.engagement, self.user, workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = profile
        config.save()
        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        html = ir_render.render_document_html(document)
        self.assertIn("Impact area", html)
        self.assertIn("Risk rating", html)
        self.assertNotIn(">Affects<", html)

    def test_plain_table_cells_use_configured_fill_color(self):
        from . import ir_render

        profile = ReportProfile.objects.create(
            name="SMOKE TEST table fill profile", empty_cell_background_color="rgba(11,22,33,1)",
        )
        css = ir_render.build_preview_css(
            type("Meta", (), {
                "body_font": "", "monospace_font": "", "bullet_character": "",
                "table_header_color": "", "empty_cell_background_color": profile.empty_cell_background_color,
            })(),
        )
        self.assertIn(".report-table td { background: rgba(11,22,33,1); }", css)

    def test_team_lead_cannot_view_profile_list(self):
        client = Client()
        login(client, make_user(User.Role.TEAM_LEAD))
        resp = client.get(reverse("report_profiles:list"))
        self.assertEqual(resp.status_code, 403)

    def test_list_page_renders_and_shows_seeded_default(self):
        resp = self.client.get(reverse("report_profiles:list"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Default")

    def test_create_profile(self):
        resp = self.client.post(
            reverse("report_profiles:list"), {"action": "create", "create-name": "SMOKE TEST profile"},
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(ReportProfile.objects.filter(name="SMOKE TEST profile").exists())

    def test_manage_text_blocks_embedded_on_profile_edit_page(self):
        profile = ReportProfile.objects.get(is_default=True)
        resp = self.client.get(reverse("report_profiles:edit", args=[profile.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Text blocks")
        self.assertContains(resp, "New text block")

    def test_text_block_create_from_profile_edit_page_redirects_back(self):
        profile = ReportProfile.objects.get(is_default=True)
        edit_url = reverse("report_profiles:edit", args=[profile.pk])
        resp = self.client.post(reverse("report_profiles:text_blocks", args=[profile.pk]), {
            "action": "create_text_block", "next": edit_url, "create-label": "SMOKE TEST inline block",
        })
        self.assertRedirects(resp, edit_url)
        block = ReportTextBlockDefinition.objects.get(label="SMOKE TEST inline block")
        self.assertEqual(block.profile, profile)

    def test_text_block_create_is_independent_per_profile(self):
        profile_a = ReportProfile.objects.create(name="SMOKE TEST profile A")
        profile_b = ReportProfile.objects.create(name="SMOKE TEST profile B")
        for profile in (profile_a, profile_b):
            resp = self.client.post(reverse("report_profiles:text_blocks", args=[profile.pk]), {
                "action": "create_text_block", "create-label": "Shared-sounding name",
            })
            self.assertEqual(resp.status_code, 302)
        blocks = ReportTextBlockDefinition.objects.filter(label="Shared-sounding name")
        self.assertEqual(blocks.count(), 2)
        self.assertEqual({b.profile_id for b in blocks}, {profile_a.pk, profile_b.pk})
        self.assertEqual({b.slug for b in blocks}, {"shared_sounding_name"})

    def test_text_block_delete_next_is_validated_against_open_redirect(self):
        profile = ReportProfile.objects.get(is_default=True)
        block = ReportTextBlockDefinition.objects.create(profile=profile, label="SMOKE TEST delete me")
        resp = self.client.post(
            reverse("report_profiles:text_block_delete", args=[block.pk]),
            {"next": "https://evil.example/"},
        )
        self.assertRedirects(resp, reverse("report_profiles:edit", args=[profile.pk]))

    def test_edit_profile_saves_template_and_block_defaults(self):
        profile = ReportProfile.objects.create(name="SMOKE TEST profile 2")
        block = ReportTextBlockDefinition.objects.create(
            profile=profile, slug="overview_block", label="Overview block",
        )
        edit_url = reverse("report_profiles:edit", args=[profile.pk])

        template_json = json.dumps({"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Overview"}]},
            {"type": "paragraph", "content": [{"type": "text", "text": f"{{{{ {block.slug} }}}}"}]},
        ]})
        resp = self.client.post(edit_url, {"template": template_json})
        self.assertEqual(resp.status_code, 302, resp.content.decode()[:2000] if resp.status_code == 200 else "")
        profile.refresh_from_db()
        self.assertIsInstance(profile.template, dict)
        self.assertEqual(profile.template["content"][0]["content"][0]["text"], "Overview")

        block_edit_url = reverse("report_profiles:profile_text_block_edit", args=[profile.pk, block.pk])
        resp_block = self.client.post(block_edit_url, {
            "content": doc_json("Custom default for this profile"),
        })
        self.assertEqual(resp_block.status_code, 302, resp_block.content.decode()[:2000] if resp_block.status_code == 200 else "")
        profile.refresh_from_db()
        self.assertIn("Custom default for this profile", profile.block_defaults[block.slug])

        resp2 = self.client.get(edit_url)
        self.assertEqual(resp2.status_code, 200)

    def test_create_profile_defaults_to_blank_template(self):
        profile = ReportProfile.objects.create(name="SMOKE TEST profile 3")
        resp = self.client.get(reverse("report_profiles:edit", args=[profile.pk]))
        self.assertEqual(resp.status_code, 200)

    def test_delete_blocked_when_default(self):
        default_profile = ReportProfile.objects.get(is_default=True)
        resp = self.client.post(reverse("report_profiles:delete", args=[default_profile.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(ReportProfile.objects.filter(pk=default_profile.pk).exists())

    def test_set_default_flips_atomically(self):
        old_default = ReportProfile.objects.get(is_default=True)
        new_profile = ReportProfile.objects.create(name="SMOKE TEST profile 4")
        resp = self.client.post(reverse("report_profiles:set_default", args=[new_profile.pk]))
        self.assertEqual(resp.status_code, 302)
        new_profile.refresh_from_db()
        old_default.refresh_from_db()
        self.assertTrue(new_profile.is_default)
        self.assertFalse(old_default.is_default)
        self.assertEqual(ReportProfile.objects.filter(is_default=True).count(), 1)

    def test_delete_non_default_profile(self):
        profile = ReportProfile.objects.create(name="SMOKE TEST profile 5")
        resp = self.client.post(reverse("report_profiles:delete", args=[profile.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(ReportProfile.objects.filter(pk=profile.pk).exists())

    def test_text_block_list_and_create(self):
        profile = ReportProfile.objects.get(is_default=True)
        resp = self.client.get(reverse("report_profiles:edit", args=[profile.pk]))
        self.assertEqual(resp.status_code, 200)
        resp = self.client.post(
            reverse("report_profiles:text_blocks", args=[profile.pk]),
            {"action": "create_text_block", "create-label": "SMOKE TEST block"},
        )
        self.assertEqual(resp.status_code, 302)
        block = ReportTextBlockDefinition.objects.get(label="SMOKE TEST block")
        self.assertEqual(block.profile, profile)

    def test_text_block_create_with_remediation_only_flag(self):
        profile = ReportProfile.objects.get(is_default=True)
        resp = self.client.post(
            reverse("report_profiles:text_blocks", args=[profile.pk]),
            {
                "action": "create_text_block", "create-label": "SMOKE TEST remediation block",
                "create-include_in_remediation_report_only": "on",
            },
        )
        self.assertEqual(resp.status_code, 302)
        block = ReportTextBlockDefinition.objects.get(label="SMOKE TEST remediation block")
        self.assertTrue(block.include_in_remediation_report_only)

    def test_configure_view_lists_profiles_and_saves_selection(self):
        custom = ReportProfile.objects.create(name="SMOKE TEST configure profile")
        resp = self.client.get(reverse("reports:configure", args=[self.engagement.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "SMOKE TEST configure profile")

        resp = self.client.post(reverse("reports:configure", args=[self.engagement.pk]), {
            "content_json": "{}", "report_profile": str(custom.pk),
        })
        self.assertEqual(resp.status_code, 302)
        config = self.assembly.get_draft_config(self.engagement)
        self.assertEqual(config.profile_id, custom.pk)

    def test_preview_shows_docx_specific_message_when_document_tab_is_empty(self):
        profile = ReportProfile.objects.create(
            name="SMOKE TEST docx-only profile", template={}, docx_template=_valid_test_template_bytes(),
        )
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = profile
        config.save()

        resp = self.client.get(reverse("reports:configure", args=[self.engagement.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "exports via an uploaded Word (.docx) template")
        self.assertNotContains(resp, "Document tab has no content configured yet")

    def test_preview_shows_generic_message_when_document_tab_empty_and_no_docx(self):
        profile = ReportProfile.objects.create(name="SMOKE TEST empty profile", template={})
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = profile
        config.save()

        resp = self.client.get(reverse("reports:configure", args=[self.engagement.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Document tab has no content configured yet")
        self.assertNotContains(resp, "exports via an uploaded Word")

    def test_ajax_preview_also_shows_empty_message(self):
        profile = ReportProfile.objects.create(
            name="SMOKE TEST docx-only ajax profile", template={}, docx_template=_valid_test_template_bytes(),
        )
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = profile
        config.save()

        resp = self.client.post(
            reverse("reports:preview", args=[self.engagement.pk]),
            data="{}", content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn("exports via an uploaded Word (.docx) template", resp.json()["html"])

    def test_toc_still_finds_entries_with_no_typed_headings_at_all(self):
        # A template with zero heading nodes of its own (some other,
        # un-headed content precedes the tag) still surfaces a heading
        # typed INSIDE a resolved text block, correctly toc_entry=True —
        # it's spliced into the same walk as the template's own headings
        # (see _parse_template's process_nodes), not lost inside whatever
        # anonymous wrapper the preceding content landed in.
        template = {"type": "doc", "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": "Some intro copy with no heading above it."}]},
            {"type": "paragraph", "content": [{"type": "text", "text": "{{ legal_notes }}"}]},
        ]}
        custom = ReportProfile.objects.create(name="SMOKE TEST no-headings profile", template=template)
        engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-NOHEAD")
        generate_project_key(engagement)
        ReportTextBlockDefinition.objects.create(profile=custom, slug="legal_notes", label="Legal notes")
        custom.block_defaults = {"legal_notes": json.dumps({"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Headless heading"}]},
        ]})}
        custom.save()
        config = self.assembly.get_draft_config(engagement)
        config.profile = custom
        config.save()

        from apps.crypto.services import get_data_key

        from . import ir_render

        project_key = get_data_key(engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=engagement, project_key=project_key, user=None,
        )
        self.assertFalse(document.sections[0].toc_entry)
        heading_section = next(s for s in document.sections if s.title == "Headless heading")
        self.assertTrue(heading_section.toc_entry)
        html = ir_render.render_toc(document)
        self.assertIn("Headless heading", html)

    def test_heading_inside_text_block_becomes_a_real_section(self):
        # A heading typed inside a text block's own content (not the
        # profile template) should get the same structural treatment as
        # one typed directly in the template: numbered, nested at the
        # tag's position, and a TOC entry for H1/H2 — see
        # assembly._walk_text_block_nodes.
        template = {"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Disclaimers"}]},
            {"type": "paragraph", "content": [{"type": "text", "text": "{{ legal_notes }}"}]},
        ]}
        custom = ReportProfile.objects.create(name="SMOKE TEST text block heading profile", template=template)
        block_json = json.dumps({"type": "doc", "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": "Intro line."}]},
            {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Data retention"}]},
            {"type": "paragraph", "content": [{"type": "text", "text": "We keep data for 90 days."}]},
        ]})
        ReportTextBlockDefinition.objects.create(
            profile=custom, slug="legal_notes", label="Legal notes",
        )
        custom.block_defaults = {"legal_notes": block_json}
        custom.save()

        config = self.assembly.get_draft_config(self.engagement)
        config.profile = custom
        config.save()

        from apps.crypto.services import get_data_key

        from . import ir_render

        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        disclaimers = next(s for s in document.sections if s.title == "Disclaimers")
        self.assertEqual(len(disclaimers.blocks), 1)
        self.assertIn("Intro line.", disclaimers.blocks[0].html)
        retention = next(c for c in disclaimers.children if c.title == "Data retention")
        self.assertTrue(retention.toc_entry)
        self.assertNotEqual(retention.number, "")
        self.assertIn("We keep data for 90 days.", retention.blocks[0].html)

        toc_html = ir_render.render_toc(document)
        self.assertIn("Data retention", toc_html)

    def test_lead_paragraph_has_no_structural_effect(self):
        # A "lead" paragraph is a size/weight bump only (report-lead CSS
        # class) — it must never gain numbering, section nesting, or a
        # table-of-contents entry the way an actual heading would.
        template = {"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Executive Summary"}]},
            {"type": "paragraph", "attrs": {"lead": True}, "content": [{"type": "text", "text": "Big intro text"}]},
        ]}
        custom = ReportProfile.objects.create(name="SMOKE TEST lead profile", template=template)
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = custom
        config.save()

        from apps.crypto.services import get_data_key

        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        self.assertEqual(len(document.sections), 1)
        section = document.sections[0]
        self.assertEqual(section.title, "Executive Summary")
        self.assertEqual(section.children, [])
        self.assertIn('class="report-lead"', section.blocks[0].html)

        from . import ir_render

        toc_html = ir_render.render_toc(document)
        self.assertNotIn("Big intro text", toc_html)

    def test_observation_and_phase_items_are_never_toc_entries(self):
        # Same reasoning as findings (see test_finding_grandchild_not_a_toc_entry):
        # a repeating content item, however many are titled, isn't a
        # chapter/subsection heading — the TOC lists only real H1/H2s.
        template = {"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Observations"}]},
            {"type": "paragraph", "content": [{"type": "text", "text": "{{ observations }}"}]},
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Testing Phases"}]},
            {"type": "paragraph", "content": [{"type": "text", "text": "{{ testing_phases }}"}]},
        ]}
        custom = ReportProfile.objects.create(name="SMOKE TEST toc entries profile", template=template)
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = custom
        config.content = {
            **config.content,
            "observations": [{"title": "Weak TLS config", "content": ""}, {"title": "", "content": ""}],
            "testing_phases": [{"title": "Reconnaissance", "content": ""}],
        }
        config.save()

        from apps.crypto.services import get_data_key

        from . import ir_render

        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        observations_section = next(s for s in document.sections if s.title == "Observations")
        self.assertFalse(observations_section.children[0].toc_entry)
        self.assertFalse(observations_section.children[1].toc_entry)

        toc_html = ir_render.render_toc(document)
        self.assertIn("Observations", toc_html)
        self.assertIn("Testing Phases", toc_html)
        self.assertNotIn("Weak TLS config", toc_html)
        self.assertNotIn("Reconnaissance", toc_html)

    def test_build_report_document_reflects_configured_profile(self):
        template = {"type": "doc", "content": [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Just findings"}]},
            {"type": "paragraph", "content": [{"type": "text", "text": "{{ finding_details }}"}]},
        ]}
        custom = ReportProfile.objects.create(name="SMOKE TEST minimal profile", template=template)
        config = self.assembly.get_draft_config(self.engagement)
        config.profile = custom
        config.content = {**config.content, "findings": {"included_ids": []}}
        config.save()

        from apps.crypto.services import get_data_key

        project_key = get_data_key(self.engagement)
        document = self.assembly.build_report_document(
            config=config, engagement=self.engagement, project_key=project_key, user=None,
        )
        self.assertEqual([s.title for s in document.sections], ["Just findings"])


class ChecklistCoverageSectionTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-CL-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        EngagementMembership.objects.create(user=self.author, engagement=self.engagement)

    def _build_document(self, *, include_tag):
        from types import SimpleNamespace

        from .assembly import build_report_document, seeded_content
        from .models import ReportProfile

        content = seeded_content(self.engagement, SimpleNamespace(content={}))
        template_content = [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Coverage"}]},
        ]
        if include_tag:
            template_content.append(
                {"type": "paragraph", "content": [{"type": "text", "text": "{{ checklist_coverage }}"}]}
            )
        profile = ReportProfile.objects.create(
            name="CHECKLIST TEST profile", template={"type": "doc", "content": template_content},
        )
        config = SimpleNamespace(content=content, is_remediation_report=False, profile=profile, profile_id=profile.id)
        return build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )

    def test_tag_absent_omits_section(self):
        document = self._build_document(include_tag=False)
        self.assertEqual(document.sections[0].blocks, [])
        self.assertEqual(document.sections[0].children, [])

    def test_tag_present_with_no_runs_shows_graceful_empty_state(self):
        document = self._build_document(include_tag=True)
        section = document.sections[0]
        self.assertEqual(section.children, [])
        html = "".join(b.html for b in section.blocks if hasattr(b, "html"))
        self.assertIn("No checklist runs", html)

    def test_tag_present_with_run_shows_items_status_and_linked_findings(self):
        finding = Finding.objects.create(
            engagement=self.engagement, title="SQL Injection", severity=Finding.Severity.HIGH,
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        run = ChecklistRun.objects.create(engagement=self.engagement, label="Web App Checklist")
        item_with_finding = ChecklistItem.objects.create(
            run=run, category="Authentication", code="AUTH-01", title="Test for SQL injection in login",
            status=ChecklistItem.Status.TESTED_FINDING_RAISED, order=1,
        )
        item_with_finding.findings.add(finding)
        ChecklistItem.objects.create(
            run=run, category="Authentication", code="AUTH-02", title="Test for account lockout",
            status=ChecklistItem.Status.TESTED_NO_FINDING, order=2,
        )
        ChecklistItem.objects.create(
            run=run, category="Session Management", code="SESS-01", title="Test session fixation",
            status=ChecklistItem.Status.NOT_APPLICABLE, order=3,
        )

        document = self._build_document(include_tag=True)
        section = document.sections[0]
        self.assertEqual(len(section.children), 1)
        run_section = section.children[0]
        self.assertEqual(run_section.title, "Web App Checklist")

        categories = {child.title: child for child in run_section.children}
        self.assertIn("Authentication", categories)
        self.assertIn("Session Management", categories)

        auth_table = categories["Authentication"].blocks[0]
        self.assertEqual(auth_table.columns, ["Code", "Item", "Status", "Linked findings"])
        rows_by_code = {row[0]: row for row in auth_table.rows}
        self.assertEqual(rows_by_code["AUTH-01"][2], "Tested — Finding Raised")
        self.assertEqual(rows_by_code["AUTH-01"][3], "F001")
        self.assertEqual(rows_by_code["AUTH-02"][2], "Tested — No Finding")
        self.assertEqual(rows_by_code["AUTH-02"][3], "—")

        session_table = categories["Session Management"].blocks[0]
        self.assertEqual(session_table.rows[0][2], "Not Applicable")


class ScanImportsSectionTests(TestCase):

    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme", reference_number="PT-SCAN-1")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        EngagementMembership.objects.create(user=self.author, engagement=self.engagement)

    def _build_document(self, *, include_tag):
        from types import SimpleNamespace

        from .assembly import build_report_document, seeded_content
        from .models import ReportProfile

        content = seeded_content(self.engagement, SimpleNamespace(content={}))
        template_content = [
            {"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "Appendix"}]},
        ]
        if include_tag:
            template_content.append(
                {"type": "paragraph", "content": [{"type": "text", "text": "{{ scan_imports }}"}]}
            )
        profile = ReportProfile.objects.create(
            name="SCAN IMPORTS TEST profile", template={"type": "doc", "content": template_content},
        )
        config = SimpleNamespace(content=content, is_remediation_report=False, profile=profile, profile_id=profile.id)
        return build_report_document(
            config=config, engagement=self.engagement,
            project_key=get_data_key(self.engagement), user=self.author,
        )

    def test_tag_absent_omits_section(self):
        document = self._build_document(include_tag=False)
        self.assertEqual(document.sections[0].blocks, [])
        self.assertEqual(document.sections[0].children, [])

    def test_tag_present_with_no_imports_shows_graceful_empty_state(self):
        document = self._build_document(include_tag=True)
        section = document.sections[0]
        html = "".join(b.html for b in section.blocks if hasattr(b, "html"))
        self.assertIn("No scanner imports", html)

    def test_tag_present_with_import_shows_tool_file_and_count(self):
        from apps.findings.models import ScanImportRecord

        ScanImportRecord.objects.create(
            engagement=self.engagement, source_format="BURP", filename="burp-export.xml",
            findings_created=5, imported_by=self.author,
        )

        document = self._build_document(include_tag=True)
        section = document.sections[0]
        table = section.blocks[0]
        self.assertEqual(table.columns, ["Tool", "File", "Imported", "Imported by", "Findings produced"])
        self.assertEqual(len(table.rows), 1)
        row = table.rows[0]
        self.assertEqual(row[0], "Burp Suite")
        self.assertEqual(row[1], "burp-export.xml")
        self.assertEqual(row[4], "5")


def _docx_bytes(build_fn) -> bytes:
    from io import BytesIO as _BytesIO

    from docx import Document as _Document

    doc = _Document()
    build_fn(doc)
    buf = _BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _valid_test_template_bytes() -> bytes:
    from docx.enum.style import WD_STYLE_TYPE

    def build(doc):
        doc.styles.add_style("Test Code Block", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_heading("{{ cover_title }}", level=1)
        doc.add_paragraph("{{ client_name }} ({{ test_type_label }})")
        doc.add_paragraph("Reviewed by: {{ reviewed_by }}  QA by: {{ qa_by }}  Approved by: {{ approved_by }}")
        doc.add_paragraph("{{p document_control }}")
        doc.add_paragraph("{{p scan_imports }}")
        doc.add_paragraph("{{p breakdown_table }}")
        doc.add_paragraph("{{p information_gathering }}")
        doc.add_paragraph("{%p for o in observations %}")
        doc.add_paragraph("{{ o.title }}")
        doc.add_paragraph("{{p o.content }}")
        doc.add_paragraph("{%p endfor %}")
        doc.add_paragraph("{%p for p in testing_phases %}")
        doc.add_paragraph("{{ p.title }}")
        doc.add_paragraph("{{p p.content }}")
        doc.add_paragraph("{%p endfor %}")
        doc.add_paragraph("{%p for finding in findings %}")
        doc.add_paragraph("{{ finding.id }} - {{ finding.title }} ({{ finding.severity }}) [{{ finding.cve_id }}]")
        doc.add_paragraph("{{p finding['technical-details'] }}")
        doc.add_paragraph("{{p finding.retest_history }}")
        doc.add_paragraph("{%p endfor %}")

    return _docx_bytes(build)


def _malformed_template_bytes() -> bytes:
    def build(doc):
        doc.add_paragraph("{%p for finding in findings %}")
        doc.add_paragraph("{{ finding.title }}")
        # Deliberately missing {%p endfor %} — malformed loop syntax.

    return _docx_bytes(build)


class TiptapToDocxTests(TestCase):
    def _tpl_with_styles(self):
        from docxtpl import DocxTemplate

        from .docx_style_roles import STYLE_ROLES

        def build(doc):
            from docx.enum.style import WD_STYLE_TYPE

            bucket_type = {
                "paragraph": WD_STYLE_TYPE.PARAGRAPH, "character": WD_STYLE_TYPE.CHARACTER,
                "table": WD_STYLE_TYPE.TABLE,
            }
            for role_key, _label, bucket in STYLE_ROLES:
                doc.styles.add_style(f"RS {role_key}", bucket_type[bucket])
            doc.add_paragraph("{{p subject }}")

        data = _docx_bytes(build)
        from io import BytesIO

        return DocxTemplate(BytesIO(data))

    def _styles_map(self):
        from .docx_style_roles import STYLE_ROLES

        return {role_key: f"RS {role_key}" for role_key, _label, _bucket in STYLE_ROLES}

    def _render(self, tiptap_json, **kwargs):
        from docx import Document

        from . import tiptap_docx

        tpl = self._tpl_with_styles()
        subdoc = tiptap_docx.tiptap_to_subdoc(
            tiptap_json, tpl=tpl, styles=self._styles_map(), image_resolver=kwargs.get("image_resolver"),
            monospace_font="Consolas",
        )
        tpl.render({"subject": subdoc})
        from io import BytesIO

        out = BytesIO()
        tpl.save(out)
        out.seek(0)
        return Document(out)

    def test_code_block_uses_mapped_style_per_line(self):
        content = json.dumps({
            "type": "doc", "content": [
                {"type": "codeBlock", "content": [{"type": "text", "text": "line1\nline2"}]},
            ],
        })
        result = self._render(content)
        styled = [p for p in result.paragraphs if p.text in ("line1", "line2")]
        self.assertEqual(len(styled), 2)
        for p in styled:
            self.assertEqual(p.style.name, "RS code_block")

    def test_bold_mark_sets_run_bold(self):
        content = json.dumps({
            "type": "doc", "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "bold", "marks": [{"type": "bold"}]}]},
            ],
        })
        result = self._render(content)
        run = next(p for p in result.paragraphs if p.text == "bold").runs[0]
        self.assertTrue(run.bold)

    def test_table_colspan_merges_cells(self):
        content = json.dumps({
            "type": "doc", "content": [
                {"type": "table", "content": [
                    {"type": "tableRow", "content": [
                        {"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "H1"}]}]},
                        {"type": "tableHeader", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "H2"}]}]},
                    ]},
                    {"type": "tableRow", "content": [
                        {"type": "tableCell", "attrs": {"colspan": 2}, "content": [
                            {"type": "paragraph", "content": [{"type": "text", "text": "spanned"}]},
                        ]},
                    ]},
                ]},
            ],
        })
        result = self._render(content)
        table = result.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells], ["H1", "H2"])
        self.assertEqual([c.text for c in table.rows[1].cells], ["spanned", "spanned"])

    def test_blank_content_produces_empty_subdoc_without_error(self):
        result = self._render("")
        # No exception, and nothing beyond the template's own static paragraph.
        self.assertEqual(len([p for p in result.paragraphs if p.text.strip()]), 0)

    def test_heading_level_5_uses_mapped_style(self):
        content = json.dumps({
            "type": "doc", "content": [
                {"type": "heading", "attrs": {"level": 5}, "content": [{"type": "text", "text": "Deep heading"}]},
            ],
        })
        result = self._render(content)
        heading = next(p for p in result.paragraphs if p.text == "Deep heading")
        self.assertEqual(heading.style.name, "RS heading_5")

    def test_oversized_rowspan_does_not_crash_and_clamps_to_actual_rows(self):
        # rowspan claims 5 rows but only 2 tableRows actually exist — a stale/corrupted
        # span (e.g. content that didn't come from the editor's own UI). Must clamp
        # instead of raising IndexError.
        content = json.dumps({
            "type": "doc", "content": [
                {"type": "table", "content": [
                    {"type": "tableRow", "content": [
                        {"type": "tableCell", "attrs": {"rowspan": 5}, "content": [
                            {"type": "paragraph", "content": [{"type": "text", "text": "spanned"}]},
                        ]},
                        {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "A"}]}]},
                    ]},
                    {"type": "tableRow", "content": [
                        {"type": "tableCell", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "B"}]}]},
                    ]},
                ]},
            ],
        })
        result = self._render(content)  # must not raise
        table = result.tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual(table.cell(0, 0).text, "spanned")
        self.assertEqual(table.cell(1, 0).text, "spanned")  # merged down to the last real row

    def test_depth_guard_stops_deeply_nested_tables(self):
        from . import tiptap_docx

        def nested_table(remaining):
            cell_content = (
                [nested_table(remaining - 1)] if remaining > 0
                else [{"type": "paragraph", "content": [{"type": "text", "text": "bottom"}]}]
            )
            return {"type": "table", "content": [
                {"type": "tableRow", "content": [{"type": "tableCell", "content": cell_content}]},
            ]}

        content = json.dumps({"type": "doc", "content": [nested_table(tiptap_docx._MAX_DEPTH + 10)]})
        result = self._render(content)  # must not raise RecursionError
        # Some outer tables render before the guard kicks in; "bottom" must never be reached.
        self.assertNotIn("bottom", [p.text for p in result.paragraphs])


class DocxStyleDiscoveryTests(TestCase):
    def test_discovers_custom_and_builtin_styles(self):
        from docx.enum.style import WD_STYLE_TYPE

        from .docx_styles import discover_styles

        def build(doc):
            doc.styles.add_style("My Code Style", WD_STYLE_TYPE.PARAGRAPH)
            doc.styles.add_style("My Char Style", WD_STYLE_TYPE.CHARACTER)
            doc.styles.add_style("My Table Style", WD_STYLE_TYPE.TABLE)

        buckets = discover_styles(_docx_bytes(build))
        self.assertIn("My Code Style", buckets["paragraph"])
        self.assertIn("Normal", buckets["paragraph"])  # built-in, always present
        self.assertIn("My Char Style", buckets["character"])
        self.assertIn("My Table Style", buckets["table"])


class DocxTextBlockPlaceholderTests(TestCase):
    def test_available_text_block_tags_reflects_this_profiles_active_blocks(self):
        from .docx_placeholders import available_text_block_tags

        profile = ReportProfile.objects.get(is_default=True)
        tags = dict(available_text_block_tags(profile))
        self.assertIn("information_gathering", tags)

    def test_unsaved_profile_returns_no_text_block_tags(self):
        from .docx_placeholders import available_text_block_tags

        self.assertEqual(available_text_block_tags(ReportProfile()), [])


class DocxTemplateUploadTests(TestCase):
    def setUp(self):
        self.profile = ReportProfile.objects.create(name="Docx profile")
        self.superadmin = make_user(User.Role.SUPERADMIN)
        self.upload_url = reverse("report_profiles:docx_template", args=[self.profile.pk])

    def test_non_docx_file_rejected(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, self.superadmin)
        upload = SimpleUploadedFile("not-a-docx.docx", b"not really a docx", content_type="application/octet-stream")
        client.post(self.upload_url, {"docx_template": upload})

        self.profile.refresh_from_db()
        self.assertFalse(self.profile.docx_template)

    def test_malformed_loop_syntax_rejected_at_upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, self.superadmin)
        upload = SimpleUploadedFile(
            "bad.docx", _malformed_template_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        client.post(self.upload_url, {"docx_template": upload})

        self.profile.refresh_from_db()
        self.assertFalse(self.profile.docx_template)

    def test_valid_template_accepted_and_redirects_to_style_map(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, self.superadmin)
        upload = SimpleUploadedFile(
            "good.docx", _valid_test_template_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        resp = client.post(self.upload_url, {"docx_template": upload})

        self.profile.refresh_from_db()
        self.assertTrue(bytes(self.profile.docx_template))
        self.assertEqual(self.profile.docx_template_filename, "good.docx")
        self.assertRedirects(resp, reverse("report_profiles:docx_style_map", args=[self.profile.pk]))

    def test_style_map_saves_selected_roles(self):
        self.profile.docx_template = _valid_test_template_bytes()
        self.profile.save(update_fields=["docx_template"])

        client = Client()
        login(client, self.superadmin)
        resp = client.post(
            reverse("report_profiles:docx_style_map", args=[self.profile.pk]),
            {"role__code_block": "Test Code Block"},
        )
        self.assertEqual(resp.status_code, 302)
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.docx_style_map.get("code_block"), "Test Code Block")

    def test_remove_clears_template_and_style_map(self):
        self.profile.docx_template = _valid_test_template_bytes()
        self.profile.docx_style_map = {"code_block": "Test Code Block"}
        self.profile.save(update_fields=["docx_template", "docx_style_map"])

        client = Client()
        login(client, self.superadmin)
        client.post(self.upload_url, {"action": "remove"})

        self.profile.refresh_from_db()
        self.assertFalse(self.profile.docx_template)
        self.assertEqual(self.profile.docx_style_map, {})

    def test_team_lead_without_permission_cannot_upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        client = Client()
        login(client, make_user(User.Role.TEAM_LEAD))
        upload = SimpleUploadedFile(
            "good.docx", _valid_test_template_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        resp = client.post(self.upload_url, {"docx_template": upload})
        self.assertEqual(resp.status_code, 403)


class DocxExportViewTests(TestCase):
    def setUp(self):
        self.engagement = Engagement.objects.create(client_name="Acme")
        generate_project_key(self.engagement)
        self.author = make_user(User.Role.CONSULTANT)
        EngagementMembership.objects.create(user=self.author, engagement=self.engagement)
        make_finding(
            self.engagement, self.author, title="Approved one",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        self.profile = ReportProfile.objects.get(is_default=True)
        self.profile.docx_template = _valid_test_template_bytes()
        self.profile.docx_style_map = {"code_block": "Test Code Block"}
        self.profile.save(update_fields=["docx_template", "docx_style_map"])

    def test_docx_export_returns_a_real_docx(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertTrue(resp.content.startswith(b"PK\x03\x04"))  # .docx is a zip container
        self.assertIn(b'attachment; filename="', resp["Content-Disposition"].encode())

    def test_docx_export_creates_log_entry(self):
        from .models import ReportExportLog

        client = Client()
        login(client, self.author)
        client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        log = ReportExportLog.objects.get(engagement=self.engagement, fmt="docx")
        self.assertEqual(log.exported_by, self.author)

    def test_docx_export_contains_finding_title(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))

        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(resp.content))
        all_text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Approved one", all_text)

    def test_export_options_hides_docx_button_when_no_template(self):
        self.profile.docx_template = None
        self.profile.save(update_fields=["docx_template"])

        client = Client()
        login(client, self.author)
        resp = client.get(reverse("reports:options", args=[self.engagement.pk]))
        self.assertNotContains(resp, 'download docx')  # sanity: no leftover copy-paste string
        self.assertFalse(resp.context["docx_available"])

    def test_download_rejected_when_no_template_configured(self):
        self.profile.docx_template = None
        self.profile.save(update_fields=["docx_template"])

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertRedirects(resp, reverse("reports:options", args=[self.engagement.pk]))

    def test_starter_template_fixture_exports_successfully(self):
        import pathlib

        fixture_path = pathlib.Path(__file__).parent / "fixtures" / "docx_starter_template.docx"
        self.profile.docx_template = fixture_path.read_bytes()
        self.profile.save(update_fields=["docx_template"])

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.startswith(b"PK\x03\x04"))

    def test_docx_export_includes_observations_and_testing_phases(self):
        from . import assembly

        config = assembly.get_draft_config(self.engagement)
        config.content = {
            **config.content,
            "observations": [{"title": "Weak TLS config", "content": doc_json("Server allows TLS 1.0.")}],
            "testing_phases": [{"title": "Reconnaissance", "content": doc_json("Passive OSINT gathering.")}],
        }
        config.save()

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(resp.content))
        all_text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Weak TLS config", all_text)
        self.assertIn("Server allows TLS 1.0.", all_text)
        self.assertIn("Reconnaissance", all_text)
        self.assertIn("Passive OSINT gathering.", all_text)

    def test_docx_export_includes_reviewer_qa_approver_names(self):
        reviewer = make_user(User.Role.SENIOR)
        qa = make_user(User.Role.SENIOR)
        approver = make_user(User.Role.TEAM_LEAD)
        self.engagement.default_reviewer = reviewer
        self.engagement.default_qa = qa
        self.engagement.default_approver = approver
        self.engagement.save(update_fields=["default_reviewer", "default_qa", "default_approver"])

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn(f"Reviewed by: {reviewer}", all_text)
        self.assertIn(f"QA by: {qa}", all_text)
        self.assertIn(f"Approved by: {approver}", all_text)

    def test_docx_export_blank_reviewer_qa_approver_when_unset(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn("Reviewed by:   QA by:   Approved by:", all_text)

    def test_docx_export_includes_retest_notes_for_remediation_report(self):
        from . import assembly

        finding = make_finding(
            self.engagement, self.author, title="Retested finding",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
        )
        self.engagement.status = Engagement.Status.REMEDIATION_TEST
        self.engagement.save(update_fields=["status"])

        client = Client()
        login(client, self.author)
        resp = client.post(
            reverse("findings:retest_create", args=[self.engagement.pk, finding.pk]),
            {"status": Finding.RetestStatus.FIXED, "notes": doc_json("Verified the patch is deployed.")},
        )
        self.assertEqual(resp.status_code, 302)

        config = assembly.get_draft_config(self.engagement)
        config.is_remediation_report = True
        config.save(update_fields=["is_remediation_report"])

        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn("Verified the patch is deployed.", all_text)
        self.assertIn("Retest notes —", all_text)

    def test_docx_export_includes_custom_text_block(self):
        self.profile.block_defaults = {
            **self.profile.block_defaults, "information_gathering": doc_json("Passive recon notes."),
        }
        self.profile.save(update_fields=["block_defaults"])

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn("Passive recon notes.", all_text)

    def test_docx_export_includes_test_type_label(self):
        self.engagement.test_type = Engagement.TestType.WEB_APPLICATION
        self.engagement.save(update_fields=["test_type"])

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn("Web Application Penetration Test", all_text)

    def test_docx_export_includes_cve_id(self):
        make_finding(
            self.engagement, self.author, title="CVE finding",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED, cve_id="CVE-2024-12345",
        )

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn("CVE-2024-12345", all_text)

    def test_docx_export_omits_observations_block_when_none_configured(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(resp.content))
        # observations/testing_phases resolve to empty lists when no entries are
        # configured — the {%p for %} loop body simply never runs, no error either.
        self.assertEqual(len(document.tables), 2)  # document_control + breakdown_table only

    def test_docx_export_breakdown_includes_chart_image(self):
        make_finding(
            self.engagement, self.author, title="Chart-triggering finding",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED, severity=Finding.Severity.HIGH,
        )

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(resp.content))
        # The breakdown chart is a picture embedded ahead of the counts
        # table, not a native Word chart object — python-docx surfaces it
        # as an inline shape.
        self.assertGreaterEqual(len(document.inline_shapes), 1)

    def test_docx_export_uses_profile_severity_and_status_label_overrides(self):
        make_finding(
            self.engagement, self.author, title="Localized severity finding",
            workflow_status=Finding.WorkflowStatus.QA_APPROVED,
            severity=Finding.Severity.CRITICAL, status=Finding.Status.OPEN,
        )
        self.profile.labels = {"severity_critical": "Critique", "status_open": "Ouverte"}
        self.profile.save(update_fields=["labels"])

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(resp.content))
        # "Ouverte" (the status_open override) only appears in the breakdown
        # table's column header — a table cell, not a top-level paragraph —
        # so paragraph text alone won't find it; include table cell text too.
        all_text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Critique", all_text)
        self.assertIn("Ouverte", all_text)
        self.assertNotIn("Critical", all_text)

    def test_docx_export_scan_imports_empty_state(self):
        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        all_text = "\n".join(p.text for p in Document(BytesIO(resp.content)).paragraphs)
        self.assertIn("No scanner imports have been recorded for this engagement.", all_text)

    def test_docx_export_includes_scan_imports_table(self):
        from apps.findings.models import ScanImportRecord

        ScanImportRecord.objects.create(
            engagement=self.engagement, source_format="NMAP", filename="scan.xml",
            findings_created=3, imported_by=self.author,
        )

        client = Client()
        login(client, self.author)
        resp = client.post(reverse("reports:download", args=[self.engagement.pk, "docx"]))
        self.assertEqual(resp.status_code, 200)

        from io import BytesIO

        from docx import Document

        document = Document(BytesIO(resp.content))
        all_text = "\n".join(p.text for p in document.paragraphs)
        self.assertNotIn("No scanner imports have been recorded", all_text)
        table_text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Nmap", table_text)
        self.assertIn("scan.xml", table_text)
        self.assertIn("3", table_text)


class SeedPentestReportFormatCommandTests(TestCase):
    """Regression coverage for `seed_pentest_report_format`. The command was
    broken outright by the reports.0022/0023 migrations (ReportTextBlockDefinition
    became profile-scoped and lost its own `content` field) — it raised a
    FieldError the moment it was run, silently, since nothing called it in
    tests or CI. This locks in the rewritten version so that regression can't
    recur unnoticed a second time."""

    def test_command_runs_without_error_and_creates_default_profile(self):
        from django.core.management import call_command

        # setUpModule()'s shared seed_report_fixtures() already made a
        # "Default" profile the instance default for the rest of this test
        # file — clear it so this test can observe its own command's
        # "claim default if nothing else is" behavior in isolation.
        ReportProfile.objects.filter(is_default=True).update(is_default=False)

        call_command("seed_pentest_report_format")

        profile = ReportProfile.objects.get(name="Default Penetration Test Report")
        self.assertTrue(profile.is_default)
        self.assertTrue(profile.template)

    def test_command_seeds_risk_assessment_finding_section_ahead_of_description(self):
        from django.core.management import call_command

        call_command("seed_pentest_report_format")

        risk = ContentSectionDefinition.objects.get(slug="risk_assessment")
        description = ContentSectionDefinition.objects.get(slug="description")
        self.assertLess(risk.order, description.order)

    def test_command_text_blocks_are_scoped_to_the_default_profile_with_real_content(self):
        from django.core.management import call_command

        call_command("seed_pentest_report_format")

        profile = ReportProfile.objects.get(name="Default Penetration Test Report")
        slugs = set(profile.text_blocks.values_list("slug", flat=True))
        self.assertIn("root_cause_analysis", slugs)
        self.assertIn("effective_security_practices", slugs)
        self.assertIn("additional_recommendations", slugs)
        self.assertIn("conclusion", slugs)
        for slug in slugs:
            self.assertTrue(profile.block_defaults.get(slug), f"{slug} has no content in block_defaults")

    def test_command_is_idempotent(self):
        from django.core.management import call_command

        call_command("seed_pentest_report_format")
        call_command("seed_pentest_report_format")

        self.assertEqual(ReportProfile.objects.filter(name="Default Penetration Test Report").count(), 1)
        self.assertEqual(ReportProfile.objects.filter(is_default=True).count(), 1)

    def test_command_does_not_steal_default_from_a_profile_an_admin_chose(self):
        """If an admin has already picked a different default, re-running this
        command later (e.g. after an upgrade) must not silently override that
        choice — only the very first run, when nothing is default yet, should
        claim it."""
        from django.core.management import call_command

        # setUpModule() already made "Default" the instance default —
        # replace it with a distinct admin-chosen one so this test's own
        # precondition ("a non-default-profile-command default exists") is
        # explicit rather than inherited.
        ReportProfile.objects.filter(is_default=True).update(is_default=False)
        other = ReportProfile.objects.create(name="Custom House Style", is_default=True)

        call_command("seed_pentest_report_format")

        other.refresh_from_db()
        self.assertTrue(other.is_default)
        seeded = ReportProfile.objects.get(name="Default Penetration Test Report")
        self.assertFalse(seeded.is_default)

    def test_command_produces_a_rendering_report_document(self):
        """End-to-end: the seeded profile's template must actually assemble
        into a real ReportDocument, not just look plausible as JSON."""
        from types import SimpleNamespace

        from django.core.management import call_command

        from .assembly import build_report_document, seeded_content

        call_command("seed_pentest_report_format")

        engagement = Engagement.objects.create(client_name="Seed Command Smoke Test")
        generate_project_key(engagement)
        profile = ReportProfile.objects.get(name="Default Penetration Test Report")
        content = seeded_content(engagement, SimpleNamespace(content={}))
        config = SimpleNamespace(content=content, is_remediation_report=False, profile=profile, profile_id=profile.id)

        document = build_report_document(
            config=config, engagement=engagement,
            project_key=get_data_key(engagement), user=None,
        )
        titles = [s.title for s in document.sections]
        self.assertIn("Executive Summary", titles)
        self.assertIn("Root Cause Analysis", titles)
        self.assertIn("Conclusion", titles)

