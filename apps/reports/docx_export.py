"""Builds a DOCX export from a superadmin-uploaded Report Profile template.

Deliberately NOT a fourth consumer of report_ir.ReportDocument — the other
three exporters (pdf_export, html_export, md_export) all render the same
generated Section/Block tree, but DOCX is a free-form Word document the
superadmin designs themselves with {{ tag }} placeholders, filled in with
docxtpl. It reads the same underlying engagement/finding data as the IR
path (reusing assembly.py's/services.py's data-fetch helpers directly) but
builds its own Jinja context instead of an IR tree.
"""
import io
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docxtpl import DocxTemplate
from jinja2.exceptions import TemplateError
from jinja2.sandbox import SandboxedEnvironment

from apps.findings.display_id import assign_display_ids
from apps.findings.models import ContentSectionDefinition, Finding

from . import assembly, services, tiptap_docx
from .labels import get_label, get_severity_label, get_status_label
from .models import ReportSettings
from .placeholders import build_placeholder_context, render_placeholders

_RGBA_RE = re.compile(r"rgba?\(\s*([\d.]+)\s*,\s*([\d.]+)\s*,\s*([\d.]+)\s*(?:,\s*[\d.]+\s*)?\)")


def sandboxed_jinja_env() -> SandboxedEnvironment:
    """The Jinja environment every .docx template render goes through.

    docxtpl uses a bare, unrestricted jinja2.Template internally when no
    jinja_env is passed to render() -- and a Report Profile's .docx
    template is uploaded content, not code RedScribe authored. Rendering
    uploaded content through an unrestricted Jinja2 environment lets a
    crafted template reach Python internals (e.g. via
    ''.__class__.__mro__) and execute arbitrary code on the server.
    SandboxedEnvironment blocks that class of attribute access while
    behaving identically for ordinary tag/loop usage, since nothing here
    registers custom filters or extensions beyond Jinja2's defaults.
    """
    return SandboxedEnvironment()


class DocxExportError(Exception):
    """Raised when a Word template can't be filled in — almost always a
    template-authoring mistake (a bad tag/loop) rather than a RedScribe
    bug, so callers should show this message directly rather than a
    generic 500."""


def _set_style_attr(obj, style_name):
    if not style_name:
        return
    try:
        obj.style = style_name
    except KeyError:
        pass


def _rgba_to_hex(value: str) -> str | None:
    if not value:
        return None
    if re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        return value.upper()
    match = _RGBA_RE.match(value.strip())
    if not match:
        return None
    r, g, b = (max(0, min(255, round(float(match.group(i))))) for i in (1, 2, 3))
    return f"{r:02X}{g:02X}{b:02X}"


def _set_cell_background(cell, rgba_or_hex: str) -> None:
    hex_color = _rgba_to_hex(rgba_or_hex)
    if not hex_color:
        return
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _apply_cell_style(cell, style_name):
    if not style_name:
        return
    for paragraph in cell.paragraphs:
        _set_style_attr(paragraph, style_name)


def _build_table(container, columns, rows, *, styles, shaded_cells=None):
    table = container.add_table(rows=len(rows) + 1, cols=len(columns))
    _set_style_attr(table, styles.get("table_normal"))
    header_style = styles.get("table_header_cell_text")
    body_style = styles.get("body_paragraph")
    for c, col in enumerate(columns):
        cell = table.cell(0, c)
        cell.text = str(col)
        _apply_cell_style(cell, header_style)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            _apply_cell_style(cell, body_style)
            if shaded_cells and (r - 1, c) in shaded_cells:
                _set_cell_background(cell, shaded_cells[(r - 1, c)])
    return table


def _simple_table_subdoc(tpl, columns, rows, *, styles, shaded_cells=None):
    subdoc = tpl.new_subdoc()
    _build_table(subdoc, columns, rows, styles=styles, shaded_cells=shaded_cells)
    return subdoc


def _document_control_subdoc(tpl, engagement, profile, styles):
    columns = [
        get_label(profile, "doc_control_col_stage"), get_label(profile, "doc_control_col_date"),
        get_label(profile, "doc_control_col_changed_by"),
    ]
    rows = assembly._document_control_rows(engagement, profile)
    return _simple_table_subdoc(tpl, columns, rows, styles=styles)


def _breakdown_chart_bars(count_rows, statuses, profile):
    """Same (label, count, color) series the IR path builds for the PDF/HTML
    bar graph (see assembly._breakdown_section) — kept in lockstep so the
    DOCX chart shows the same bars, in the same order, as every other
    format."""
    open_label = get_status_label(profile, "OPEN")
    closed_label = get_status_label(profile, "CLOSED")
    bars = []
    for row in count_rows:
        if Finding.Status.OPEN in statuses:
            bars.append((f"{row['label']} ({open_label})", row["open_count"], row["open_color"]))
        if Finding.Status.CLOSED in statuses and row["closed_count"]:
            bars.append((f"{row['label']} ({closed_label})", row["closed_count"], row["closed_color"]))
    return bars


def _draw_centered_label(draw, label, cx, top_y, font, fill):
    # Split "Critical (Open)" into two lines so it doesn't overrun a
    # neighboring bar's slot at the narrow widths a chart embedded in a
    # Word page tends to get.
    if " (" in label and label.endswith(")"):
        main, suffix = label.split(" (", 1)
        lines = [main, "(" + suffix]
    else:
        lines = [label]
    y = top_y
    line_height = getattr(font, "size", 20) + 4
    for line in lines:
        w = draw.textlength(line, font=font)
        draw.text((cx - w / 2, y), line, fill=fill, font=font)
        y += line_height


def _breakdown_chart_png(count_rows, statuses, profile) -> bytes:
    """Renders the same severity/status breakdown the PDF/HTML bar graph
    shows as a raster image, since python-docx has no charting API to draw
    a native Word chart object with. Pillow is already a dependency (TOTP
    QR codes), so this needs no new package."""
    from PIL import Image, ImageDraw, ImageFont

    bars = _breakdown_chart_bars(count_rows, statuses, profile)
    if not bars:
        return b""

    scale = 3  # supersampled, then displayed scaled down — crisp at Word's rendered size
    width, height = 900 * scale, 420 * scale
    margin_left, margin_right = 50 * scale, 30 * scale
    margin_top, margin_bottom = 40 * scale, 90 * scale
    plot_w = width - margin_left - margin_right
    plot_h = height - margin_top - margin_bottom
    baseline_y = margin_top + plot_h

    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    label_font = ImageFont.load_default(size=20 * scale)
    value_font = ImageFont.load_default(size=22 * scale)

    draw.line([(margin_left, baseline_y), (width - margin_right, baseline_y)], fill="#C3C2B7", width=scale)

    max_value = max((count for _label, count, _color in bars), default=0) or 1
    slot_w = plot_w / len(bars)
    bar_w = min(slot_w * 0.55, 110 * scale)

    for i, (label, value, color) in enumerate(bars):
        hex_color = _rgba_to_hex(color) or "888888"
        rgb = tuple(int(hex_color[j:j + 2], 16) for j in (0, 2, 4))
        bar_h = (value / max_value) * plot_h if value else 0
        x0 = margin_left + i * slot_w + (slot_w - bar_w) / 2
        y0 = baseline_y - bar_h
        x1 = x0 + bar_w
        if bar_h:
            draw.rectangle([x0, y0, x1, baseline_y], fill=rgb)
        value_text = str(value)
        tw = draw.textlength(value_text, font=value_font)
        draw.text((x0 + bar_w / 2 - tw / 2, y0 - 30 * scale), value_text, fill="#131A22", font=value_font)
        _draw_centered_label(draw, label, x0 + bar_w / 2, baseline_y + 12 * scale, label_font, "#52514E")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _breakdown_table_subdoc(tpl, findings, profile, options, styles):
    _filtered, _severities, statuses, count_rows = assembly._breakdown_counts(findings, profile, options)
    subdoc = tpl.new_subdoc()

    chart_bytes = _breakdown_chart_png(count_rows, statuses, profile)
    if chart_bytes:
        run = subdoc.add_paragraph().add_run()
        width_emu = tiptap_docx.content_width_emu(tpl)
        try:
            if width_emu:
                run.add_picture(io.BytesIO(chart_bytes), width=width_emu)
            else:
                run.add_picture(io.BytesIO(chart_bytes), width=Inches(6))
        except Exception:
            # A malformed page-size/template edge case shouldn't sink the
            # whole export — same "best effort" contract as _firm_logo_subdoc.
            pass

    columns = [
        get_label(profile, "vuln_table_col_severity"),
        get_status_label(profile, "OPEN"), get_status_label(profile, "CLOSED"),
    ]
    rows = [[row["label"], str(row["open_count"]), str(row["closed_count"])] for row in count_rows]
    shaded = {}
    for i, row in enumerate(count_rows):
        shaded[(i, 1)] = row["open_color"]
        shaded[(i, 2)] = row["closed_color"]
    _build_table(subdoc, columns, rows, styles=styles, shaded_cells=shaded)
    return subdoc


def _assessment_team_subdoc(tpl, engagement, profile, styles):
    subdoc = tpl.new_subdoc()
    data = assembly._assessment_team_data(engagement)
    if not data:
        _set_style_attr(subdoc.add_paragraph("No assessment team members recorded."), styles.get("body_paragraph"))
        return subdoc
    for entry in data:
        member = entry["member"]
        _set_style_attr(subdoc.add_paragraph(f"{member} — {entry['role_name']}"), styles.get("heading_3"))
        if entry["background"]:
            _set_style_attr(subdoc.add_paragraph(entry["background"]), styles.get("body_paragraph"))
        for qualification in entry["qualifications_list"] or []:
            _set_style_attr(subdoc.add_paragraph(qualification), styles.get("bullet_list"))
    return subdoc


def _checklist_coverage_subdoc(tpl, engagement, findings, profile, styles):
    subdoc = tpl.new_subdoc()
    display_id_lookup = {f.id: f.display_id for f in findings}
    data = assembly._checklist_coverage_data(engagement, display_id_lookup)
    if not data:
        _set_style_attr(
            subdoc.add_paragraph("No checklist runs have been recorded for this engagement."),
            styles.get("body_paragraph"),
        )
        return subdoc
    columns = [
        get_label(profile, "checklist_col_code"), get_label(profile, "checklist_col_item"),
        get_label(profile, "checklist_col_status"), get_label(profile, "checklist_col_linked_findings"),
    ]
    for run_data in data:
        _set_style_attr(subdoc.add_paragraph(run_data["run_label"]), styles.get("heading_2"))
        for category in run_data["categories"]:
            _set_style_attr(subdoc.add_paragraph(category["category"]), styles.get("heading_3"))
            _build_table(subdoc, columns, category["rows"], styles=styles)
    return subdoc


def _scan_imports_subdoc(tpl, engagement, profile, styles):
    subdoc = tpl.new_subdoc()
    data = assembly._scan_imports_data(engagement)
    if not data:
        _set_style_attr(
            subdoc.add_paragraph("No scanner imports have been recorded for this engagement."),
            styles.get("body_paragraph"),
        )
        return subdoc
    columns = [
        get_label(profile, "scan_imports_col_tool"), get_label(profile, "scan_imports_col_filename"),
        get_label(profile, "scan_imports_col_date"), get_label(profile, "scan_imports_col_imported_by"),
        get_label(profile, "scan_imports_col_findings"),
    ]
    rows = [
        [row["tool"], row["filename"], row["imported_at"], row["imported_by"], str(row["findings_created"])]
        for row in data
    ]
    _build_table(subdoc, columns, rows, styles=styles)
    return subdoc


def _firm_logo_subdoc(tpl, settings_obj):
    subdoc = tpl.new_subdoc()
    run = subdoc.add_paragraph().add_run()
    try:
        run.add_picture(io.BytesIO(bytes(settings_obj.firm_logo)), width=Inches(1.5))
    except Exception:
        pass
    return subdoc


def _finding_context(
    tpl, finding, *, project_key, image_resolver, styles, active_sections, is_remediation_report,
    monospace_font, text_transform, profile,
):
    section_json = services._decrypt_finding_sections(finding, project_key, active_sections)
    entry = {
        "id": finding.display_id,
        "title": finding.title,
        "severity": get_severity_label(profile, finding.severity),
        "status": get_status_label(profile, finding.status),
        "cvss_score": finding.cvss_score or "—",
        "cvss_vector": finding.cvss_vector or "—",
        "cve_id": finding.cve_id or "—",
        "classifications": [str(tag) for tag in finding.classifications.all()],
        "affects": list(finding.affects_list),
    }
    for definition in active_sections:
        if definition.include_in_remediation_report_only and not is_remediation_report:
            continue
        content = section_json.get(definition.slug)
        entry[definition.slug] = (
            tiptap_docx.tiptap_to_subdoc(
                content, tpl=tpl, styles=styles, image_resolver=image_resolver,
                text_transform=text_transform, monospace_font=monospace_font,
            )
            if content else tpl.new_subdoc()
        )

    entry["retest_history"] = tpl.new_subdoc()
    if is_remediation_report:
        retest_records = services._decrypt_retest_notes(finding, project_key)
        if retest_records:
            columns = [
                get_label(profile, "retest_col_date"), get_label(profile, "retest_col_tested_by"),
                get_label(profile, "retest_col_result"),
            ]
            rows = [[r["date"], r["tested_by"], r["status_display"]] for r in retest_records]
            _build_table(entry["retest_history"], columns, rows, styles=styles)
            # Each record's own note text — matches assembly.py's PDF/HTML/MD path (a
            # "Retest notes — <date> (<status>)" subsection per record with notes). Was
            # dropped in the original DOCX implementation as a scope trim; added back for
            # content parity with the other formats.
            width = tiptap_docx.content_width_emu(tpl)
            for record in retest_records:
                if not record["notes_json"]:
                    continue
                heading = f"Retest notes — {record['date']} ({record['status_display']})"
                _set_style_attr(entry["retest_history"].add_paragraph(heading), styles.get("heading_3"))
                tiptap_docx.render_tiptap_into(
                    entry["retest_history"], record["notes_json"], styles=styles, image_resolver=image_resolver,
                    text_transform=text_transform, monospace_font=monospace_font, content_width_emu=width,
                )
    return entry


def _entry_list(tpl, entries, styles, *, image_resolver, text_transform, monospace_font):
    """Builds the list of {"title": <plain str>, "content": <subdoc>} dicts that
    observations/testing_phases expose to the template as a real loop variable — the
    template author controls layout entirely themselves (heading level, a table row,
    whatever), the same way the finding loop works, rather than RedScribe imposing a
    fixed internal layout the way a single {{p tag }} would."""
    result = []
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        title = entry.get("title", "")
        result.append({
            "title": text_transform(title) if text_transform else title,
            "content": tiptap_docx.tiptap_to_subdoc(
                entry.get("content", ""), tpl=tpl, styles=styles, image_resolver=image_resolver,
                text_transform=text_transform, monospace_font=monospace_font,
            ),
        })
    return result


def _text_block_context(tpl, profile, content, styles, *, is_remediation_report, **common_body_kwargs):
    """One {{p slug }} rich-content tag per this profile's own active custom text
    blocks (ReportTextBlockDefinition — e.g. "information_gathering", "disclaimers")
    — the same tags the Document-tab/PDF/HTML/MD path resolves via
    assembly._resolve_custom_tag_nodes. Missing from DOCX entirely until now, so a
    template referencing one silently rendered blank (Jinja's default-undefined
    behavior for a context key that was never supplied)."""
    result = {}
    for definition in profile.text_blocks.filter(is_active=True):
        if definition.include_in_remediation_report_only and not is_remediation_report:
            continue
        raw = assembly.dynamic_tag_value(content, definition.slug, profile)
        result[definition.slug] = (
            tiptap_docx.tiptap_to_subdoc(raw, tpl=tpl, styles=styles, **common_body_kwargs)
            if raw else tpl.new_subdoc()
        )
    return result


def _base_context(tpl, *, profile, settings_obj, engagement, findings, content, styles, user, image_resolver, is_remediation_report):
    placeholder_context = build_placeholder_context(engagement=engagement, user=user)

    def substitute(text):
        return render_placeholders(text, placeholder_context)

    common_body_kwargs = dict(image_resolver=image_resolver, text_transform=substitute, monospace_font=profile.monospace_font)

    context = dict(placeholder_context)
    context.update(_text_block_context(
        tpl, profile, content, styles, is_remediation_report=is_remediation_report, **common_body_kwargs,
    ))
    context.update({
        "cover_title": substitute(profile.cover_title),
        "classification_label": profile.classification_label,
        "firm_name": settings_obj.firm_name,
        "test_type_label": engagement.get_test_type_display() if engagement.test_type else "",
        "reviewed_by": str(engagement.default_reviewer) if engagement.default_reviewer_id else "",
        "qa_by": str(engagement.default_qa) if engagement.default_qa_id else "",
        "approved_by": str(engagement.default_approver) if engagement.default_approver_id else "",
        "document_control": _document_control_subdoc(tpl, engagement, profile, styles),
        "assessment_team": _assessment_team_subdoc(tpl, engagement, profile, styles),
        "checklist_coverage": _checklist_coverage_subdoc(tpl, engagement, findings, profile, styles),
        "scan_imports": _scan_imports_subdoc(tpl, engagement, profile, styles),
        "breakdown_table": _breakdown_table_subdoc(
            tpl, findings, profile, content.get("breakdown") if isinstance(content, dict) else None, styles,
        ),
        "observations": _entry_list(
            tpl, assembly._dict_rows(content.get("observations")), styles, **common_body_kwargs,
        ),
        "testing_phases": _entry_list(
            tpl, assembly._dict_rows(content.get("testing_phases")), styles, **common_body_kwargs,
        ),
    })
    if settings_obj.firm_logo:
        context["firm_logo"] = _firm_logo_subdoc(tpl, settings_obj)
    else:
        context["firm_logo"] = tpl.new_subdoc()
    return context, substitute


def build_docx(*, config, engagement, project_key, user) -> bytes:
    profile = assembly.get_effective_profile(config)
    if profile is None or not profile.docx_template:
        raise DocxExportError("This report profile has no Word (.docx) template configured.")

    tpl = DocxTemplate(io.BytesIO(bytes(profile.docx_template)))
    styles = profile.docx_style_map if isinstance(profile.docx_style_map, dict) else {}
    content = assembly.sanitize_content(config.content)
    findings = assembly._select_findings(engagement, content.get("findings", {}))
    assign_display_ids(findings, prefix=profile.finding_id_prefix or "F")

    settings_obj = ReportSettings.get_solo()
    image_resolver = services.make_docx_image_resolver(engagement, project_key)
    context, substitute = _base_context(
        tpl, profile=profile, settings_obj=settings_obj, engagement=engagement, findings=findings,
        content=content, styles=styles, user=user, image_resolver=image_resolver,
        is_remediation_report=config.is_remediation_report,
    )

    active_sections = list(ContentSectionDefinition.objects.filter(is_active=True))
    context["findings"] = [
        _finding_context(
            tpl, f, project_key=project_key, image_resolver=image_resolver, styles=styles,
            active_sections=active_sections, is_remediation_report=config.is_remediation_report,
            monospace_font=profile.monospace_font, text_transform=substitute, profile=profile,
        )
        for f in findings
    ]

    try:
        tpl.render(context, jinja_env=sandboxed_jinja_env())
    except TemplateError as exc:
        raise DocxExportError(f"This Word template couldn't be rendered: {exc}") from exc

    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()


def build_dummy_context(profile):
    """Cheap, DB-light context (no real engagement) used to dry-run a newly
    uploaded template at upload time — see docx_template_validation.py.
    Includes exactly one dummy finding so a malformed {%p for %}/{%tr for %}
    loop body is actually exercised, not skipped over an empty list."""
    tpl = DocxTemplate(io.BytesIO(bytes(profile.docx_template)))
    styles = profile.docx_style_map if isinstance(profile.docx_style_map, dict) else {}
    active_sections = list(ContentSectionDefinition.objects.filter(is_active=True))

    context = {
        "client_name": "Example Client", "reference_number": "PROJ-0001", "scope": "example.com",
        "start_date": "", "end_date": "", "report_date": "", "prepared_by": "",
        "reviewed_by": "", "qa_by": "", "approved_by": "", "test_type_label": "",
        "cover_title": profile.cover_title, "classification_label": profile.classification_label,
        "firm_name": "Example Firm", "document_control": tpl.new_subdoc(), "assessment_team": tpl.new_subdoc(),
        "checklist_coverage": tpl.new_subdoc(), "scan_imports": tpl.new_subdoc(),
        "breakdown_table": tpl.new_subdoc(), "firm_logo": tpl.new_subdoc(),
        "observations": [{"title": "Example observation", "content": tpl.new_subdoc()}],
        "testing_phases": [{"title": "Example phase", "content": tpl.new_subdoc()}],
    }
    for definition in profile.text_blocks.filter(is_active=True):
        context[definition.slug] = tpl.new_subdoc()
    dummy_finding = {
        "id": "F001", "title": "Example finding", "severity": "High", "status": "Open",
        "cvss_score": "7.5", "cvss_vector": "—", "cve_id": "—", "classifications": [], "affects": [],
        "retest_history": tpl.new_subdoc(),
    }
    for definition in active_sections:
        dummy_finding[definition.slug] = tpl.new_subdoc()
    context["findings"] = [dummy_finding]
    return tpl, context
