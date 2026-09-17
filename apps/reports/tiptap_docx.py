"""Renders a finding's TipTap/ProseMirror JSON into a docxtpl subdocument,
applying named Word styles from the uploaded template's own style catalog
(apps.reports.docx_style_roles / docx_styles) rather than styles RedScribe
invents — this is the fix for the failure mode that killed the prior DOCX
export attempt (rich content like code blocks/images had nowhere to pick up
the template's own packaged styling).

Sibling to tiptap_render.py — the node types handled here (text, paragraph,
heading, bulletList/orderedList/listItem, blockquote, codeBlock,
horizontalRule, pageBreak, hardBreak, image, table/tableRow/tableHeader/
tableCell) and marks (bold, italic, underline, strike, code, highlight,
link) must stay in parity with that file's node-type switch.
"""
import io
import json

from docx.enum.text import WD_BREAK
from docx.image.image import Image as _DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run

_ALLOWED_LINK_PROTOCOLS = ("http://", "https://", "mailto:")
_MAX_DEPTH = 60


def tiptap_to_subdoc(doc_json, *, tpl, styles, image_resolver=None, text_transform=None, monospace_font=""):
    """doc_json: decrypted raw TipTap JSON string (NOT pre-flattened HTML).
    styles: profile.docx_style_map dict, role_key -> style name (or missing/blank).
    image_resolver: callable(src) -> (bytes, content_type) | None — raw bytes,
    unlike services._make_image_resolver's data-URI string for HTML/PDF.
    """
    subdoc = tpl.new_subdoc()
    render_tiptap_into(
        subdoc, doc_json, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
        monospace_font=monospace_font, content_width_emu=content_width_emu(tpl),
    )
    return subdoc


def render_tiptap_into(container, doc_json, *, styles, image_resolver=None, text_transform=None, monospace_font="", content_width_emu=None):
    """Renders TipTap JSON into an existing container instead of creating a new subdoc —
    used by tiptap_to_subdoc above (one field, one subdoc), and directly by callers that
    need to combine several entries into a single subdoc (e.g. docx_export.py's
    observations/testing-phases lists: one heading + one rendered body per entry, all in
    the same subdoc, the same way block_registry.py's _observations/_testing_phases build
    one Section per entry for the other formats)."""
    parsed = None
    if doc_json:
        try:
            parsed = json.loads(doc_json)
        except (TypeError, ValueError):
            parsed = None
    if isinstance(parsed, dict):
        for node in parsed.get("content") or []:
            _render_block_node(
                container, node, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
                content_width_emu=content_width_emu, monospace_font=monospace_font,
            )


def content_width_emu(tpl):
    try:
        section = tpl.docx.sections[0]
        width = section.page_width - section.left_margin - section.right_margin
        return int(width) if width else None
    except Exception:
        return None


def _set_style(set_style_fn, style_name):
    if not style_name:
        return
    try:
        set_style_fn(style_name)
    except KeyError:
        pass


def _render_block_node(
    container, node, *, styles, image_resolver, text_transform, content_width_emu, monospace_font,
    depth=0, list_style_override=None,
):
    if not isinstance(node, dict) or depth > _MAX_DEPTH:
        return
    node_type = node.get("type")
    content = node.get("content") or []
    attrs = node.get("attrs") or {}

    if node_type == "paragraph":
        style_name = list_style_override or styles.get("body_paragraph")
        paragraph = container.add_paragraph()
        _set_style(lambda name: setattr(paragraph, "style", name), style_name)
        _render_inline(
            paragraph, content, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
            content_width_emu=content_width_emu, monospace_font=monospace_font,
        )
        return

    if node_type == "heading":
        level = attrs.get("level")
        level = level if level in (1, 2, 3, 4, 5, 6) else 3
        style_name = styles.get(f"heading_{level}")
        paragraph = container.add_paragraph()
        _set_style(lambda name: setattr(paragraph, "style", name), style_name)
        _render_inline(
            paragraph, content, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
            content_width_emu=content_width_emu, monospace_font=monospace_font,
        )
        return

    if node_type in ("bulletList", "orderedList"):
        item_style = styles.get("bullet_list" if node_type == "bulletList" else "numbered_list")
        for child in content:
            if isinstance(child, dict) and child.get("type") == "listItem":
                _render_list_item(
                    container, child, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
                    content_width_emu=content_width_emu, monospace_font=monospace_font,
                    depth=depth + 1, item_style=item_style,
                )
        return

    if node_type == "blockquote":
        style_name = styles.get("blockquote")
        for child in content:
            _render_block_node(
                container, child, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
                content_width_emu=content_width_emu, monospace_font=monospace_font,
                depth=depth + 1, list_style_override=style_name,
            )
        return

    if node_type == "codeBlock":
        style_name = styles.get("code_block")
        text = "".join(c.get("text", "") for c in content if isinstance(c, dict) and c.get("type") == "text")
        if text_transform:
            text = text_transform(text)
        for line in (text.split("\n") if text else [""]):
            paragraph = container.add_paragraph(line)
            _set_style(lambda name: setattr(paragraph, "style", name), style_name)
        return

    if node_type == "horizontalRule":
        container.add_paragraph("—" * 20)
        return

    if node_type == "pageBreak":
        container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        return

    if node_type == "image":
        _add_image_block(
            container, attrs, styles=styles, image_resolver=image_resolver, content_width_emu=content_width_emu,
        )
        return

    if node_type == "table":
        _render_table(
            container, node, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
            content_width_emu=content_width_emu, monospace_font=monospace_font, depth=depth + 1,
        )
        return

    # Unknown/unsupported block node type — best-effort: render its plain
    # text content rather than silently dropping it (mirrors tiptap_render's
    # fallback of returning children_html for a node type it doesn't know).
    text_bits: list[str] = []
    _collect_text(node, text_bits)
    if text_bits:
        joined = "".join(text_bits)
        paragraph = container.add_paragraph(text_transform(joined) if text_transform else joined)
        _set_style(lambda name: setattr(paragraph, "style", name), styles.get("body_paragraph"))


def _render_list_item(container, node, *, styles, image_resolver, text_transform, content_width_emu, monospace_font, depth, item_style):
    for child in node.get("content") or []:
        if not isinstance(child, dict):
            continue
        if child.get("type") == "paragraph":
            _render_block_node(
                container, child, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
                content_width_emu=content_width_emu, monospace_font=monospace_font,
                depth=depth, list_style_override=item_style,
            )
        else:
            # Nested list / other block inside a list item — flattened to
            # the same indent level in v1 (python-docx list nesting needs
            # numPr/ilvl XML, not parameterizable per depth by a named
            # style alone). Known limitation, not a silent drop.
            _render_block_node(
                container, child, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
                content_width_emu=content_width_emu, monospace_font=monospace_font, depth=depth,
            )


def _collect_text(node, out: list[str]):
    if not isinstance(node, dict):
        return
    if node.get("type") == "text":
        out.append(node.get("text", ""))
    for child in node.get("content") or []:
        _collect_text(child, out)


def _render_inline(paragraph, nodes, *, styles, image_resolver, text_transform, content_width_emu, monospace_font):
    for node in nodes or []:
        if not isinstance(node, dict):
            continue
        node_type = node.get("type")
        if node_type == "text":
            text = node.get("text", "")
            if text_transform:
                text = text_transform(text)
            _add_text_run(paragraph, text, node.get("marks"), styles=styles, monospace_font=monospace_font)
        elif node_type == "hardBreak":
            paragraph.add_run().add_break()
        elif node_type == "image":
            _add_image_block(
                paragraph, node.get("attrs") or {}, styles=styles, image_resolver=image_resolver,
                content_width_emu=content_width_emu, inline_paragraph=paragraph,
            )
        # Other inline-position node types are silently skipped — mirrors
        # tiptap_render._render_node's fallback for unrecognized types.


def _add_text_run(paragraph, text, marks, *, styles, monospace_font):
    marks = marks or []
    href = None
    for mark in marks:
        if isinstance(mark, dict) and mark.get("type") == "link":
            candidate = ((mark.get("attrs") or {}).get("href") or "").strip()
            if candidate.lower().startswith(_ALLOWED_LINK_PROTOCOLS):
                href = candidate

    run = _add_hyperlink_run(paragraph, text, href) if href else paragraph.add_run(text)

    for mark in marks:
        mark_type = mark.get("type") if isinstance(mark, dict) else None
        if mark_type == "bold":
            run.bold = True
        elif mark_type == "italic":
            run.italic = True
        elif mark_type == "underline":
            run.underline = True
        elif mark_type == "strike":
            run.font.strike = True
        elif mark_type == "code":
            style_name = styles.get("inline_code")
            if style_name:
                try:
                    run.style = style_name
                    continue
                except KeyError:
                    pass
            if monospace_font:
                run.font.name = monospace_font
        elif mark_type == "highlight":
            from docx.enum.text import WD_COLOR_INDEX

            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _add_hyperlink_run(paragraph, text, url):
    # python-docx has no first-class hyperlink API — this is the standard
    # raw-oxml recipe: a real relationship + a <w:hyperlink> wrapping a run.
    # "Hyperlink" is one of Word's built-in latent styles, valid to
    # reference even if not explicitly present in the template's styles.xml.
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run_element = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    style_ref = OxmlElement("w:rStyle")
    style_ref.set(qn("w:val"), "Hyperlink")
    run_props.append(style_ref)
    run_element.append(run_props)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run_element.append(text_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    return Run(run_element, paragraph)


def _picture_width(image_bytes, content_width_emu):
    if not content_width_emu:
        return None
    try:
        natural_width = _DocxImage.from_blob(image_bytes).width
    except Exception:
        return None
    return content_width_emu if natural_width > content_width_emu else None


def _add_image_block(container, attrs, *, styles, image_resolver, content_width_emu, inline_paragraph=None):
    if not image_resolver:
        return
    resolved = image_resolver(attrs.get("src"))
    if not resolved:
        return
    image_bytes, _content_type = resolved
    paragraph = inline_paragraph if inline_paragraph is not None else container.add_paragraph()
    run = paragraph.add_run()
    width = _picture_width(image_bytes, content_width_emu)
    try:
        if width:
            run.add_picture(io.BytesIO(image_bytes), width=width)
        else:
            run.add_picture(io.BytesIO(image_bytes))
    except Exception:
        return
    alt = (attrs.get("alt") or "").strip()
    if alt and inline_paragraph is None:
        caption = container.add_paragraph(alt)
        _set_style(lambda name: setattr(caption, "style", name), styles.get("image_caption"))


def _compute_table_grid(rows):
    occupied: dict[tuple[int, int], bool] = {}
    placements = []
    max_col = 0
    for r, row in enumerate(rows):
        c = 0
        for cell in row.get("content") or []:
            if not isinstance(cell, dict):
                continue
            while occupied.get((r, c)):
                c += 1
            cell_attrs = cell.get("attrs") or {}
            colspan = cell_attrs.get("colspan")
            colspan = colspan if isinstance(colspan, int) and colspan > 0 else 1
            rowspan = cell_attrs.get("rowspan")
            rowspan = rowspan if isinstance(rowspan, int) and rowspan > 0 else 1
            placements.append((r, c, rowspan, colspan, cell))
            for rr in range(r, r + rowspan):
                for cc in range(c, c + colspan):
                    occupied[(rr, cc)] = True
            max_col = max(max_col, c + colspan)
            c += colspan
    return len(rows), max_col, placements


def _render_table(container, node, *, styles, image_resolver, text_transform, content_width_emu, monospace_font, depth=0):
    if depth > _MAX_DEPTH:
        return
    rows = [c for c in (node.get("content") or []) if isinstance(c, dict) and c.get("type") == "tableRow"]
    if not rows:
        return
    n_rows, n_cols, placements = _compute_table_grid(rows)
    if n_cols == 0:
        return
    table = container.add_table(rows=n_rows, cols=n_cols)
    _set_style(lambda name: setattr(table, "style", name), styles.get("table_normal"))
    has_real_header = any(
        isinstance(cell, dict) and cell.get("type") == "tableHeader"
        for row in rows for cell in (row.get("content") or [])
    )
    for r, c, rowspan, colspan, cell in placements:
        is_header = cell.get("type") == "tableHeader" or (not has_real_header and r == 0)
        default_style = styles.get("table_header_cell_text") if is_header else styles.get("body_paragraph")
        target = table.cell(r, c)
        _render_cell_content(
            target, cell.get("content") or [], styles=styles, image_resolver=image_resolver,
            text_transform=text_transform, content_width_emu=content_width_emu, monospace_font=monospace_font,
            default_style=default_style, depth=depth,
        )
        if rowspan > 1 or colspan > 1:
            # Clamp to the table's actual bounds rather than trusting rowspan/colspan
            # outright — a stale or corrupted span (e.g. content that didn't originate
            # from the editor's own UI) extending past the real row/column count would
            # otherwise crash table.cell() with an IndexError instead of just rendering
            # a smaller merge than requested.
            merge_r = min(r + rowspan - 1, n_rows - 1)
            merge_c = min(c + colspan - 1, n_cols - 1)
            if merge_r > r or merge_c > c:
                target.merge(table.cell(merge_r, merge_c))


def _render_cell_content(cell, nodes, *, styles, image_resolver, text_transform, content_width_emu, monospace_font, default_style, depth=0):
    first = True
    for node in nodes or []:
        if not isinstance(node, dict):
            continue
        if first and node.get("type") == "paragraph":
            # Reuse the cell's existing default empty paragraph for the
            # first block instead of leaving it blank and adding another.
            paragraph = cell.paragraphs[0]
            _set_style(lambda name: setattr(paragraph, "style", name), default_style)
            _render_inline(
                paragraph, node.get("content") or [], styles=styles, image_resolver=image_resolver,
                text_transform=text_transform, content_width_emu=content_width_emu, monospace_font=monospace_font,
            )
            first = False
            continue
        first = False
        override = default_style if node.get("type") == "paragraph" else None
        _render_block_node(
            cell, node, styles=styles, image_resolver=image_resolver, text_transform=text_transform,
            content_width_emu=content_width_emu, monospace_font=monospace_font, list_style_override=override,
            depth=depth,
        )
